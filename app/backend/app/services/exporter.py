"""Excel 导出（API-F14-01、CMP-F14-02，AC-F14-01、NFR-003）。

- 每个视图复用页面接口的同一服务函数（reports.py / budget_table），因此范围过滤与页面一致，导出天然只含用户范围内的行。
- 表头与页面当前显示列一致：渠道×站点偏差在 BI 未到时不导出「收入差额拆分」与费率三列（AC-F10-05）。
- 金额写入数值并保留 2 位小数（ROUND_HALF_UP，与页面显示相同的舍入）；比率写入数值并按页面小数位设置百分比格式。
- 达成条列以 Excel 数据条呈现达成率（不显示数值），与页面的达成条列对应。
"""
from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.calc.periods import fmt_month
from app.core.errors import Forbidden, NotFound, ValidationFailed
from app.core.money import CENT, HUNDRED, fmt_foreign, round_to
from app.core.scope import ScopeContext
from app.models import BudgetTask, BudgetVersion, Category, FillingUnit, PlatformUnit, Store, TargetAdjustment
from app.models.enums import UnitType
from app.services import reports as R

AMOUNT_FMT = "#,##0.00"
INT_FMT = "#,##0"
HEADER_FILL = PatternFill("solid", fgColor="EEF1F5")
BOLD = Font(bold=True)


# ---------------------------------------------------------------------------
# 单元格值
# ---------------------------------------------------------------------------
@dataclass
class Cell:
    value: Any
    fmt: str | None = None
    bar: bool = False   # 达成条列（数据条）


def amount(v: str | None) -> Cell:
    """金额：2 位小数数值；空值「—」。"""
    if v is None:
        return Cell("—")
    return Cell(Decimal(v).quantize(CENT, rounding="ROUND_HALF_UP"), AMOUNT_FMT)


def integer(v: str | None) -> Cell:
    if v is None:
        return Cell("—")
    return Cell(round_to(Decimal(v), 0), INT_FMT)


def percent(v: str | None, places: int = 2) -> Cell:
    """百分数字符串 → 小数数值 + 百分比格式（显示位数与页面一致）。"""
    if v is None:
        return Cell("—")
    ratio = round_to(Decimal(v), places) / HUNDRED
    return Cell(ratio, "0." + "0" * places + "%" if places else "0%")


def signed_percent(v: str | None, places: int = 1) -> Cell:
    if v is None:
        return Cell("—")
    ratio = round_to(Decimal(v), places) / HUNDRED
    p = "0." + "0" * places
    return Cell(ratio, f"+{p}%;−{p}%;{p}%")


def number(v: str | None, places: int = 2) -> Cell:
    if v is None:
        return Cell("—")
    p = "0." + "0" * places if places else "0"
    return Cell(round_to(Decimal(v), places), f"#,##{p}")


def bar(v: str | None) -> Cell:
    """达成条：数据条的值 = 达成率（0–120% 映射到条宽）；无法计算时留空。"""
    if v is None:
        return Cell(None, bar=True)
    return Cell(min(max(Decimal(v), Decimal(0)), Decimal(120)) / HUNDRED, "0%", bar=True)


def text(v: Any) -> Cell:
    return Cell("—" if v is None or v == "" else v)


def note_text(note: dict | None) -> Cell:
    if not note:
        return Cell("—")
    return Cell(f"{note['text']}（{note['author']} · {note['updated_at']}）")


# ---------------------------------------------------------------------------
# 工作表
# ---------------------------------------------------------------------------
@dataclass
class Sheet:
    title: str
    groups: list[tuple[str, int]] | None   # 第一行分组表头（标题, 跨列数）；None 表示单行表头
    headers: list[str]
    rows: list[list[Cell]] = field(default_factory=list)
    total_rows: set[int] = field(default_factory=set)


def _write(ws: Worksheet, sheet: Sheet) -> None:
    r = 1
    if sheet.groups:
        c = 1
        for title, span in sheet.groups:
            ws.cell(row=1, column=c, value=title)
            if span > 1:
                ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + span - 1)
            c += span
        r = 2
    for i, h in enumerate(sheet.headers, 1):
        ws.cell(row=r, column=i, value=h)
    for hr in range(1, r + 1):
        for i in range(1, len(sheet.headers) + 1):
            cell = ws.cell(row=hr, column=i)
            cell.font, cell.fill = BOLD, HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
    bar_cols: set[int] = set()
    for ri, row in enumerate(sheet.rows):
        excel_row = r + 1 + ri
        for ci, cell in enumerate(row, 1):
            target = ws.cell(row=excel_row, column=ci, value=cell.value)
            if cell.fmt:
                target.number_format = cell.fmt
            if cell.bar:
                bar_cols.add(ci)
            if ri in sheet.total_rows:
                target.font = BOLD
    last = r + len(sheet.rows)
    for ci in bar_cols:
        letter = get_column_letter(ci)
        if last > r:
            ws.conditional_formatting.add(
                f"{letter}{r + 1}:{letter}{last}",
                DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color="5B6B7F",
                            showValue=False))
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    for ci, h in enumerate(sheet.headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(10, min(40, len(h) * 2 + 4))


def build_workbook(sheets: list[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        _write(wb.create_sheet(sh.title[:31]), sh)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 视图：公司汇总（PAGE-F10-01 §8.3，三个工作表）
# ---------------------------------------------------------------------------
def company_sheets(db: Session, ctx: ScopeContext, filters: dict) -> tuple[str, list[Sheet]]:
    data = R.company_summary(db, ctx, _int(filters.get("year")))
    if data["empty"]:
        raise ValidationFailed(data["message"])
    ach = data["achievement"]
    np_sheet = Sheet("净利润", None, ["月份", "年初原始目标净利润", "最新调整后目标净利润", "实际净利润", "达成率（原始）",
                                   "达成率（调整后）", "达成条（调整后）"])
    for m in ach["months"]:
        np_sheet.rows.append([text(m["month"]), amount(m["original_target"]), amount(m["adjusted_target"]),
                              amount(m["actual"]), percent(m["rate_original"], 1), percent(m["rate_adjusted"], 1),
                              bar(m["rate_adjusted"])])
    np_sheet.total_rows.add(len(np_sheet.rows))
    np_sheet.rows.append([text(f"1–{len(ach['months'])} 月累计"), amount(ach["original"]["target_sum"]),
                          amount(ach["adjusted"]["target_sum"]), amount(ach["actual_sum"]),
                          percent(ach["original"]["rate"], 1), percent(ach["adjusted"]["rate"], 1),
                          bar(ach["adjusted"]["rate"])])
    g = data["growth"]
    if g["available"]:
        gs = Sheet("增速对比", None, ["项目", f"{g['prior_period']} 实际", f"{g['period']} 实际", "增额", "增速", "明细"])
        gs.rows.append([text("收入合计"), amount(g["revenue"]["prior"]), amount(g["revenue"]["current"]),
                        amount(g["revenue"]["increase"]), signed_percent(g["revenue"]["growth"]), text("渠道×站点偏差")])
        gs.rows.append([text("部门费用合计"), amount(g["dept_expense"]["prior"]), amount(g["dept_expense"]["current"]),
                        amount(g["dept_expense"]["increase"]), signed_percent(g["dept_expense"]["growth"]),
                        text("部门费用报告")])
    else:
        gs = Sheet("增速对比", None, ["说明"], [[text(g["message"])]])
    acc = Sheet("预测准确度", None, ["月份", "所属季度快照", "快照预测净利润率", "实际净利润率", "偏差（百分点）", "判定"])
    for a in data["accuracy"]:
        if a["actual_imported"]:
            acc.rows.append([text(a["month"]), text(a["snapshot"]), percent(a["forecast_margin"]),
                             percent(a["actual_margin"]), number(a["deviation_pp"]),
                             text("超阈值" if a["flagged"] else "未超")])
        else:
            acc.rows.append([text(a["month"]), text(a["snapshot"]), percent(a["forecast_margin"]), text("实际未导入"),
                             text("—"), text("待金蝶实际")])
    return f"公司汇总_截至{data['through']}.xlsx", [np_sheet, gs, acc]


# ---------------------------------------------------------------------------
# 视图：渠道×站点偏差（PAGE-F10-02 §8，主表 + 展开明细）
# ---------------------------------------------------------------------------
def channel_main_columns(bi_available: bool) -> tuple[list[tuple[str, int]], list[str]]:
    """主表列（与页面 PAGE-F10-02 §5.2 一致）；BI 未到时整组隐藏拆分列。"""
    groups = [("店铺", 1), ("渠道", 1), ("站点", 1), ("收入", 6)]
    headers = ["店铺", "渠道", "站点", "预算", "实际", "实际（本币）", "差额", "达成率", "达成条"]
    if bi_available:
        groups.append(("收入差额拆分", 4))
        headers += ["量差", "价差", "汇率差", "口径差"]
    groups += [("毛利", 4), ("毛利率", 3), ("偏差说明", 1)]
    headers += ["预算", "实际", "差额", "达成率", "预算", "实际", "差（百分点）", "偏差说明"]
    return groups, headers


def channel_detail_columns(bi_available: bool) -> list[str]:
    headers = ["店铺 / 客户", "项目", "预算", "实际", "差额", "达成率", "达成条"]
    if bi_available:
        headers += ["预算费率", "实际费率", "费率差"]
    return headers


def _channel_row(row: dict, bi: bool) -> list[Cell]:
    kind = row["kind"]
    rev, gp, gm = row["revenue"], row["gross_profit"], row["gross_margin"]
    if kind == "store":
        local = fmt_foreign(Decimal(row["revenue_local_actual"]), row["currency"]) \
            if bi and row["revenue_local_actual"] is not None else ("待 BI" if not bi else "—")
    else:
        local = "—"
    cells = [text(row["name"]), text(row["channel"]), text(row["site"]), amount(rev["budget"]), amount(rev["actual"]),
             text(local), amount(rev["diff"]), percent(rev["rate"]), bar(rev["rate"])]
    if bi:
        sp = row.get("split")
        if kind == "b2b":
            cells += [text("B 端不做量价拆分"), text("—"), text("—"), text("—")]
        elif sp is None:
            cells += [text("—")] * 4
        else:
            cells += [amount(sp["volume"]), amount(sp["price"]),
                      text(sp["missing_fx_text"]) if sp["missing_fx"] else amount(sp["fx"]),
                      text("—") if sp["missing_fx"] else amount(sp["basis"])]
    cells += [amount(gp["budget"]), amount(gp["actual"]), amount(gp["diff"]), percent(gp["rate"]),
              percent(gm["budget"]), percent(gm["actual"]), number(gm["diff_pp"]), note_text(row.get("note"))]
    return cells


def channel_sheets(db: Session, ctx: ScopeContext, filters: dict) -> tuple[str, list[Sheet]]:
    data = R.channel_site(db, ctx, filters.get("month") or None, filters.get("baseline") or "adj",
                          filters.get("channel") or None)
    if data["empty"]:
        raise ValidationFailed(data["message"])
    bi = data["bi_available"]
    groups, headers = channel_main_columns(bi)
    main = Sheet("渠道站点偏差", groups, headers)
    detail = Sheet("明细", None, channel_detail_columns(bi))
    all_rows = data["rows"] + data["b2b_rows"] + ([data["unassigned"]] if data["unassigned"] else [])
    for row in data["rows"]:
        main.rows.append(_channel_row(row, bi))
    if data["b2b_rows"]:
        main.rows.append([text("B 端（按客户）")] + [Cell(None)] * (len(headers) - 1))
        for row in data["b2b_rows"]:
            main.rows.append(_channel_row(row, bi))
    if data["unassigned"]:
        u = data["unassigned"]
        cells = [text("未归属渠道"), text("—"), text("—")] + [text("—")] * 5 + [Cell(None, bar=True)]
        if bi:
            cells += [text("—")] * 4
        cells += [text("—"), amount(u["gross_profit"]["actual"]), text("—"), text("—"), text("—"), text("—"),
                  text("—"), note_text(u.get("note"))]
        main.rows.append(cells)
    t = data["total"]
    tcells = [text(t["label"]), Cell(None), Cell(None), amount(t["revenue"]["budget"]), amount(t["revenue"]["actual"]),
              text("—"), amount(t["revenue"]["diff"]), percent(t["revenue"]["rate"]), bar(t["revenue"]["rate"])]
    if bi:
        sp = t["split"]
        if sp is None:
            tcells += [text("—")] * 4
        else:
            tcells += [amount(sp["volume"]), amount(sp["price"]),
                       text(sp["missing_fx_text"]) if sp["missing_fx"] else amount(sp["fx"]),
                       text("—") if sp["missing_fx"] else amount(sp["basis"])]
    tcells += [amount(t["gross_profit"]["budget"]), amount(t["gross_profit"]["actual"]),
               amount(t["gross_profit"]["diff"]), percent(t["gross_profit"]["rate"]),
               percent(t["gross_margin"]["budget"]), percent(t["gross_margin"]["actual"]),
               number(t["gross_margin"]["diff_pp"]), Cell(None)]
    main.total_rows.add(len(main.rows))
    main.rows.append(tcells)

    for row in all_rows:
        name = row["name"]
        if row["kind"] != "unassigned":
            rev = row["revenue"]
            cells = [text(name), text("收入"), amount(rev["budget"]), amount(rev["actual"]), amount(rev["diff"]),
                     percent(rev["rate"]), bar(rev["rate"])]
            detail.rows.append(cells + ([text("—")] * 3 if bi else []))
        for line in row["lines"]:
            if row["kind"] == "unassigned":
                cells = [text(name), text(line["label"]), text("—"), amount(line["actual"]), text("—"), text("—"),
                         Cell(None, bar=True)]
                detail.rows.append(cells + ([text("—")] * 3 if bi else []))
                continue
            if Decimal(line["budget"] or 0) == 0 and Decimal(line["actual"] or 0) == 0:
                continue
            cells = [text(name), text(line["label"]), amount(line["budget"]), amount(line["actual"]),
                     amount(line["diff"]), percent(line["rate"]), bar(line["rate"])]
            if bi:
                cells += [percent(line["budget_rate"], 1), percent(line["actual_rate"], 1), amount(line["rate_diff"])]
            detail.rows.append(cells)
        gp = row["gross_profit"]
        if row["kind"] == "unassigned":
            cells = [text(name), text("毛利影响"), text("—"), amount(gp["actual"]), text("—"), text("—"),
                     Cell(None, bar=True)]
        else:
            cells = [text(name), text("毛利"), amount(gp["budget"]), amount(gp["actual"]), amount(gp["diff"]),
                     percent(gp["rate"]), bar(gp["rate"])]
        detail.rows.append(cells + ([text("—")] * 3 if bi else []))
    base_name = R.BASELINES[data["baseline"]]
    return f"渠道站点偏差_{data['month']}_{base_name}.xlsx", [main, detail]


# ---------------------------------------------------------------------------
# 视图：品类×新老品偏差（PAGE-F10-03 §8）
# ---------------------------------------------------------------------------
CATEGORY_GROUPS = [("品类", 1), ("新品/老品", 1), ("SPU 数", 3), ("销量（件）", 3), ("品效（折后 · 人民币元/SPU）", 3),
                   ("收入（元）", 3), ("偏差说明", 1)]
CATEGORY_HEADERS = ["品类", "新品/老品"] + ["预算", "实际", "差额"] * 4 + ["偏差说明"]


def _category_row(r: dict, note: bool) -> list[Cell]:
    return [text(r["category"]), text(r["kind"]),
            integer(r["spu"]["budget"]), integer(r["spu"]["actual"]), integer(r["spu"]["diff"]),
            integer(r["qty"]["budget"]), integer(r["qty"]["actual"]), integer(r["qty"]["diff"]),
            amount(r["efficiency"]["budget"]), amount(r["efficiency"]["actual"]), amount(r["efficiency"]["diff"]),
            amount(r["revenue"]["budget"]), amount(r["revenue"]["actual"]), amount(r["revenue"]["diff"]),
            note_text(r.get("note")) if note else Cell(None)]


def category_sheets(db: Session, ctx: ScopeContext, filters: dict) -> tuple[str, list[Sheet]]:
    data = R.category_report(db, ctx, filters.get("month") or None)
    if data["empty"]:
        raise ValidationFailed(data["message"])
    sh = Sheet("品类新老品偏差", CATEGORY_GROUPS, CATEGORY_HEADERS)
    for g in data["groups"]:
        for r in g["rows"]:
            sh.rows.append(_category_row(r, True))
        sh.total_rows.add(len(sh.rows))
        sh.rows.append(_category_row(g["subtotal"], False))
    sh.total_rows.add(len(sh.rows))
    sh.rows.append(_category_row(data["total"], False))
    return f"品类新老品偏差_{data['month']}.xlsx", [sh]


# ---------------------------------------------------------------------------
# 视图：部门费用报告（PAGE-F10-04 §8）
# ---------------------------------------------------------------------------
DEPT_HEADERS = ["部门", "预算科目", "全年预算", "累计实际", "额度使用率", "时间进度", "达成条", "偏差说明"]


def dept_sheets(db: Session, ctx: ScopeContext, filters: dict) -> tuple[str, list[Sheet]]:
    data = R.dept_report(db, ctx, _int(filters.get("year")), _int(filters.get("department_id")))
    if data["empty"]:
        raise ValidationFailed(data["message"])
    sh = Sheet("部门费用报告", None, DEPT_HEADERS)
    progress = data["time_progress"]

    def line(r, note: bool):
        return [text(r["department"]), text(r["account"]), amount(r["budget"]), amount(r["actual"]),
                percent(r["usage"]), percent(progress), bar(r["usage"]),
                note_text(r.get("note")) if note else Cell(None)]

    for g in data["groups"]:
        sh.total_rows.add(len(sh.rows))
        sh.rows.append(line(g["subtotal"], False))
        for r in g["rows"]:
            sh.rows.append(line(r, True))
    sh.total_rows.add(len(sh.rows))
    sh.rows.append(line(data["total"], False))
    return f"部门费用报告_截至{data['through']}.xlsx", [sh]


# ---------------------------------------------------------------------------
# 视图：平台渠道预算表（PAGE-F04-02 填报表格：品类 × 新老品 × 项目，1–12 月 + 全年）
# ---------------------------------------------------------------------------
PLATFORM_INPUTS = [("spu_count", "SPU 数量", "int"), ("avg_sales", "单 SPU 平均销量", "num1"),
                   ("avg_price", "平均单价", "amount"), ("discount_rate", "折扣率", "ratio"),
                   ("refund_rate", "退款率", "ratio"), ("ad_rate", "广告费率", "ratio"),
                   ("commission_rate", "佣金率", "ratio")]
PLATFORM_CALCS = [("sales_qty", "销量", "num1"), ("revenue_local", "收入（本币）", "amount"),
                  ("revenue_cny", "收入（人民币）", "amount"), ("gross_profit", "毛利（人民币）", "amount"),
                  ("gross_margin", "毛利率", "margin")]


def _budget_cell(kind: str, v: Decimal | None) -> Cell:
    if v is None:
        return Cell("—")
    if kind == "int":
        return Cell(int(v), INT_FMT)
    if kind == "num1":
        return Cell(round_to(v, 1), "#,##0.0")
    if kind == "amount":
        return Cell(round_to(v, 2), AMOUNT_FMT)
    return Cell(round_to(v, 3), "0.0%")  # ratio / margin：小数，百分比 1 位


def budget_platform_sheets(db: Session, ctx: ScopeContext, filters: dict) -> tuple[str, list[Sheet]]:
    unit_id = _int(filters.get("unit_id"))
    unit = db.get(FillingUnit, unit_id) if unit_id else None
    if unit is None or unit.unit_type != UnitType.PLATFORM:
        raise NotFound("填报单位不存在")
    store = db.get(Store, unit.scope_object_id)
    if store is None:
        raise NotFound("填报单位不存在")
    if ctx.logistics_only_fields:
        raise Forbidden()
    ctx.ensure_store(store.id, store.channel_id)
    version = db.get(BudgetVersion, unit.version_id)
    year = version.year
    units = db.scalars(select(PlatformUnit).where(PlatformUnit.version_id == unit.version_id,
                                                  PlatformUnit.store_id == store.id)).all()
    cats = {c.id: c for c in db.scalars(select(Category)).all()}
    months = [f"{year}-{m:02d}" for m in range(1, 13)]
    headers = ["品类", "新/老品", "项目"] + [f"{m} 月" for m in range(1, 13)] + ["全年"]
    sh = Sheet("预算表", None, headers)
    by_group: dict[tuple[int, bool], dict[str, PlatformUnit]] = {}
    for u in units:
        by_group.setdefault((u.category_id, u.is_new), {})[fmt_month(u.month)] = u
    keys = sorted(by_group, key=lambda k: (cats[k[0]].sort, cats[k[0]].id, not k[1]))
    totals: dict[str, dict[str, Decimal]] = {m: {"sales_qty": Decimal(0), "revenue_local": Decimal(0),
                                                 "revenue_cny": Decimal(0), "gross_profit": Decimal(0)}
                                             for m in months}
    for key in keys:
        cat, is_new = cats[key[0]].name, "新品" if key[1] else "老品"
        cells_by_month = by_group[key]
        for field_name, label, kind in PLATFORM_INPUTS + PLATFORM_CALCS:
            lab = f"{label}（{store.currency}）" if field_name == "avg_price" else label
            row = [text(cat), text(is_new), text(lab)]
            year_total = None
            for m in months:
                u = cells_by_month.get(m)
                v = getattr(u, field_name) if u is not None else None
                if field_name == "gross_margin":
                    row.append(percent_margin(v))
                else:
                    row.append(_budget_cell(kind, Decimal(v) if v is not None else None))
                if field_name in ("sales_qty", "revenue_local", "revenue_cny", "gross_profit") and v is not None:
                    totals[m][field_name] += v
                    year_total = (year_total or Decimal(0)) + v
            if field_name in ("sales_qty", "revenue_local", "revenue_cny", "gross_profit"):
                row.append(_budget_cell(kind, year_total))
            elif field_name == "gross_margin":
                rev = sum((cells_by_month[m].revenue_cny or 0 for m in cells_by_month), Decimal(0))
                gp = sum((cells_by_month[m].gross_profit or 0 for m in cells_by_month), Decimal(0))
                row.append(percent_margin(gp / rev if rev else None))
            else:
                row.append(text("—"))
            sh.rows.append(row)
    for field_name, label in (("sales_qty", "销量"), ("revenue_local", "收入（本币）"), ("revenue_cny", "收入（人民币）"),
                              ("gross_profit", "毛利（人民币）"), ("gross_margin", "毛利率")):
        row = [text("店铺合计"), text("—"), text(label)]
        for m in months:
            t = totals[m]
            if field_name == "gross_margin":
                row.append(percent_margin(t["gross_profit"] / t["revenue_cny"] if t["revenue_cny"] else None))
            else:
                row.append(_budget_cell("num1" if field_name == "sales_qty" else "amount", t[field_name]))
        if field_name == "gross_margin":
            rev = sum((totals[m]["revenue_cny"] for m in months), Decimal(0))
            gp = sum((totals[m]["gross_profit"] for m in months), Decimal(0))
            row.append(percent_margin(gp / rev if rev else None))
        else:
            row.append(_budget_cell("num1" if field_name == "sales_qty" else "amount",
                                    sum((totals[m][field_name] for m in months), Decimal(0))))
        sh.total_rows.add(len(sh.rows))
        sh.rows.append(row)
    task_name = _unit_task_name(db, unit)
    return f"预算表_{store.name}_{task_name}.xlsx", [sh]


def percent_margin(v: Decimal | None) -> Cell:
    if v is None:
        return Cell("—")
    return Cell(round_to(Decimal(v), 3), "0.0%")


def _unit_task_name(db: Session, unit: FillingUnit) -> str:
    if unit.task_id:
        return db.scalar(select(BudgetTask.name).where(BudgetTask.id == unit.task_id)) or ""
    adj = db.get(TargetAdjustment, unit.adjustment_id)
    return f"{adj.year}年第{adj.quarter}季度目标调整" if adj else ""


# ---------------------------------------------------------------------------
# 分发
# ---------------------------------------------------------------------------
VIEWS: dict[str, Callable[[Session, ScopeContext, dict], tuple[str, list[Sheet]]]] = {
    "report-company": company_sheets,
    "report-channel": channel_sheets,
    "report-category": category_sheets,
    "report-dept": dept_sheets,
    "budget-platform": budget_platform_sheets,
}


def export(db: Session, ctx: ScopeContext, view: str, filters: dict) -> tuple[str, bytes]:
    builder = VIEWS.get(view)
    if builder is None:
        raise ValidationFailed("不支持导出该页面", details={"view": view})
    name, sheets = builder(db, ctx, filters)
    return name, build_workbook(sheets)


def _int(v) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(v)
    except (TypeError, ValueError) as e:
        raise ValidationFailed("筛选条件不合法") from e


__all__ = ["export", "VIEWS", "channel_main_columns", "channel_detail_columns", "CATEGORY_HEADERS", "DEPT_HEADERS"]
