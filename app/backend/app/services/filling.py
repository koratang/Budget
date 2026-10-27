"""填报服务 CMP-F04-02（REQ-F04-01~05、REQ-F06-01/02、REQ-F13-01；SEQ-F04-01；DEC-005~007、DEC-023）。

表格结构（app/backend/README.md「填报表格结构」）：按月布局，列 key 为 'YYYY-MM'；行 key 编码「单元维度|输入项」：
  pu:<店铺>:<站点>:<品类>:<1新/0老>|<驱动>        平台渠道单元（platform_unit）
  b2b:<客户>|<输入项>                              B 端单元（b2b_unit，B 端业务的六项）
  b2bc:<客户>|<比率>                               B 端客户的采购成本占比（财务BP）/ 物流四项（物流物控）
  ccr:<渠道>:<站点>:<品类>|<比率>                  渠道品类比例（channel_category_ratio）
  csr:<渠道>:<站点>|<比率>                         渠道站点比例（库存索赔率、其他运营费用率）
  opl|<项目>                                       公司层其他损益（company_other_pl）
  de:<预算科目>|amount                             部门费用（dept_expense）
计算行与合计行由服务端给出，客户端夹带的派生字段一律忽略（RULE-01）；合计用 SQL SUM 聚合未舍入值（DEC-005/006）。
值为字符串：输入格以无损的展示格式下发（可作为基准值回传），计算格与合计格为 2 位小数的展示值。
"""
from __future__ import annotations

import importlib
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select, update
from sqlalchemy.orm import Session

from app.calc.formulas import company_net_profit, exceeds_ramp_limit, gross_margin, new_channel_month_index
from app.calc.periods import fmt_month, parse_month, year_months
from app.calc.recalc import recalc_b2b_units, recalc_platform_units
from app.core.audit import write_audit
from app.core.auth import CurrentUser
from app.core.clock import clock
from app.core.errors import Conflict, Forbidden, NotFound, ValidationFailed
from app.core.money import (MINUS, MoneyError, fmt_amount, fmt_cny, fmt_foreign, fmt_percent, round_to,
                            to_decimal)
from app.core.scope import ScopeContext
from app.models import (B2bUnit, BCustomer, BudgetAccount, BudgetTask, BudgetVersion, Category, Channel,
                        ChannelCategoryRatio, ChannelSiteRatio, CompanyOtherPl, Department, DeptExpense,
                        FillingUnit, FxBudget, ImportBatch, PlatformUnit, Reserve, RuleVersion, Site, Store,
                        TargetAdjustment)
from app.models.enums import (AdjustmentStatus, ImportSource, ImportStatus, OtherPlItem, Role, RuleVersionStatus,
                              UnitStatus, UnitType, VersionType)
from app.services import tasking as T

ZERO = Decimal(0)
HUNDRED = Decimal(100)
OVER_LIMIT_LABEL = "超出新渠道放量规则"

# ---------------------------------------------------------------------------
# 输入项与取值范围（AC-F04-02、技术方案 §4.3「输入校验」与「金额边界」）
# ---------------------------------------------------------------------------
RANGES = {
    "int": "允许范围：大于等于 0 的整数",
    "nonneg": "允许范围：大于等于 0",
    "price": "允许范围：大于等于 0，最多 2 位小数",
    "rate": "允许范围：0% 到 100%",
    "money2": "允许范围：大于等于 0，最多 2 位小数",
    "signed2": "允许范围：任意金额（损失或支出填负数），最多 2 位小数",
}
MAX_PLACES = {"int": 0, "nonneg": 4, "price": 2, "rate": 4, "money2": 2, "signed2": 2}

PLATFORM_INPUTS = {"spu_count": ("SPU 数量", "int"), "avg_sales": ("单 SPU 平均销量", "nonneg"),
                   "avg_price": ("平均单价（{cur}）", "price"), "discount_rate": ("折扣率", "rate"),
                   "refund_rate": ("退款率", "rate"), "ad_rate": ("广告费率", "rate"),
                   "commission_rate": ("佣金率", "rate")}
PLATFORM_CALCS = {"sales_qty": "销量", "revenue_local": "收入（{cur}）", "revenue_cny": "收入（人民币）",
                  "gross_profit": "毛利（人民币）", "gross_margin": "毛利率"}
B2B_INPUTS = {"sales_local": ("销售额（{cur}）", "money2"), "refund_rate": ("退款率", "rate"),
              "promotion_rate": ("推广费用率", "rate"), "other_opex_rate": ("其他运营费用率", "rate"),
              "commission_rate": ("平台佣金率", "rate"), "inventory_claim_rate": ("库存索赔率", "rate")}
B2B_CALCS = {"revenue_cny": "收入（人民币）", "gross_profit": "毛利（人民币）", "gross_margin": "毛利率"}
PURCHASE = {"purchase_cost_rate": ("采购成本占比", "rate")}
LOGISTICS = {"first_leg_rate": ("头程费率", "rate"), "last_mile_rate": ("尾程费率", "rate"),
             "storage_rate": ("仓储费率", "rate"), "tariff_rate": ("关税率", "rate")}
SITE_RATES = {"inventory_claim_rate": ("库存索赔率", "rate"), "other_opex_rate": ("其他运营费用率", "rate")}
OTHER_PL = {k: (v, "signed2") for k, v in OtherPlItem.LABELS.items()}
DEPT = {"amount": ("预算", "money2")}

# 各单位类型可写的「行前缀 → 输入项」
WRITABLE: dict[str, dict[str, dict]] = {
    UnitType.PLATFORM: {"pu": PLATFORM_INPUTS},
    UnitType.B2B: {"b2b": B2B_INPUTS},
    UnitType.COST_RATIO: {"ccr": PURCHASE, "b2bc": PURCHASE, "csr": SITE_RATES, "opl": OTHER_PL},
    UnitType.LOGISTICS_RATIO: {"ccr": LOGISTICS, "b2bc": LOGISTICS},
    UnitType.DEPT_EXPENSE: {"de": DEPT},
}
# 所有角色的输入项（不属于本单位的输入项视为越权，派生字段视为伪造并忽略）
ALL_INPUT_ITEMS = (set(PLATFORM_INPUTS) | set(B2B_INPUTS) | set(PURCHASE) | set(LOGISTICS) | set(SITE_RATES)
                   | set(OTHER_PL) | set(DEPT))

TABS = {
    UnitType.PLATFORM: ("main",),
    UnitType.B2B: ("main",),
    UnitType.COST_RATIO: ("platform", "b2b", "site", "other"),
    UnitType.LOGISTICS_RATIO: ("platform", "b2b"),
    UnitType.DEPT_EXPENSE: ("main",),
}


class CellError(Exception):
    def __init__(self, raw: Any, reason: str, kind: str):
        self.raw = "" if raw is None else str(raw)
        self.reason = reason
        self.kind = kind
        super().__init__(reason)

    @property
    def grid_message(self) -> str:
        if self.reason == "不是数字":
            return f"请输入数字。{RANGES[self.kind]}"
        return f"输入「{self.raw}」未保存，{RANGES[self.kind]}"

    @property
    def template_message(self) -> str:
        return f"「{self.raw}」{self.reason}，{RANGES[self.kind]}"


def _places(d: Decimal) -> int:
    exp = d.as_tuple().exponent
    return -exp if isinstance(exp, int) and exp < 0 else 0


def parse_cell(kind: str, raw: Any) -> Decimal | int:
    """原始输入 → 存储值（比率转小数、SPU 数量转整数）。清空按 0 保存（DEC-F04-02）。越界抛 CellError。"""
    if isinstance(raw, float):  # Excel 数值单元格：按文本口径转十进制，不进入金额计算路径
        raw = format(raw, ".15g")
    s = "" if raw is None else str(raw).strip()
    for ch in (",", "，", " ", "$", "¥", "￥"):
        s = s.replace(ch, "")
    if kind == "rate" and s.endswith("%"):
        s = s[:-1]
    if s == "":
        d = ZERO
    else:
        try:
            d = to_decimal(s)
        except MoneyError:
            raise CellError(raw, "不是数字", kind) from None
    if kind == "int":
        if d != d.to_integral_value():
            raise CellError(raw, "不是整数", kind)
        if d < 0:
            raise CellError(raw, "小于 0", kind)
        return int(d)
    if kind == "rate":
        if d < 0 or d > HUNDRED:
            raise CellError(raw, "超出取值范围", kind)
    elif kind != "signed2" and d < 0:
        raise CellError(raw, "小于 0", kind)
    if _places(d) > MAX_PLACES[kind]:
        raise CellError(raw, "小数位数过多", kind)
    return d / HUNDRED if kind == "rate" else d


def _lossless(v: Decimal, places: int) -> str:
    q = round_to(v, places)
    s = f"{abs(q):,.{places}f}" if q == v else format(abs(v), ",f")
    return (MINUS if v < 0 else "") + s


def fmt_input(kind: str, v) -> str:
    """输入格展示值（无损，可回传作基准值）：整数千分位、比率百分数 2 位、金额 2 位。"""
    if v is None:
        v = ZERO
    if kind == "int":
        return f"{int(v):,}"
    v = Decimal(v)
    if kind == "rate":
        return _lossless(v * HUNDRED, 2) + "%"
    if kind == "nonneg":
        return _lossless(v, 1)
    return _lossless(v, 2)


def fmt_calc(item: str, v) -> str:
    if item == "gross_margin":
        return fmt_percent(v, 2)
    if v is None:
        return "—"
    if item == "sales_qty":  # 销量（件）统一按整数千分位显示，与原型一致
        q = round_to(Decimal(v), 0)
        return (MINUS if q < 0 else "") + f"{abs(q):,.0f}"
    return fmt_amount(v)


def margin_of(gp: Decimal | None, rev: Decimal | None) -> Decimal | None:
    if gp is None or rev is None:
        return None
    return gross_margin(gp, rev)


def _sum(values) -> Decimal | None:
    vals = [v for v in values if v is not None]
    return sum(vals, ZERO) if vals else None


# ---------------------------------------------------------------------------
# 列定义
# ---------------------------------------------------------------------------
def _dim(key: str, label: str, width: int = 120) -> dict:
    return {"key": key, "label": label, "kind": "dim", "type": "text", "editable": False, "width": width}


def _month_cols(year: int, editable: bool, *, total_label: str | None = "全年") -> list[dict]:
    cols = [{"key": fmt_month(m), "label": f"{m.month} 月", "kind": "value", "type": "number",
             "editable": editable, "group": f"{year}"} for m in year_months(year)]
    if total_label:
        cols.append({"key": "year_total", "label": total_label, "kind": "value", "type": "number",
                     "editable": False, "width": 130})
    return cols


def _version(db: Session, unit: FillingUnit) -> BudgetVersion:
    return db.get(BudgetVersion, unit.version_id)


def _can_write(db: Session, ctx: ScopeContext, unit: FillingUnit) -> bool:
    try:
        T.ensure_writable(db, ctx, unit)
    except (Conflict, Forbidden):
        return False
    return True


# ---------------------------------------------------------------------------
# 新渠道放量（DEC-023、AC-F13-02）
# ---------------------------------------------------------------------------
def ramp_rule(db: Session) -> RuleVersion | None:
    return db.scalar(select(RuleVersion).where(RuleVersion.status == RuleVersionStatus.EFFECTIVE,
                                               RuleVersion.content["kind"].astext == "new_channel_ramp")
                     .order_by(RuleVersion.id.desc()).limit(1))


def ramp_limits(rv: RuleVersion | None) -> list[Decimal] | None:
    if rv is None:
        return None
    return [to_decimal(str(x)) for x in (rv.content or {}).get("limits", [])]


def new_channel_start(db: Session, store: Store) -> date | None:
    if store.is_new_channel and store.new_channel_start:
        return store.new_channel_start
    channel = db.get(Channel, store.channel_id)
    if channel is not None and channel.is_new_channel and channel.new_channel_start:
        return channel.new_channel_start
    return None


def _store_month_revenue(db: Session, version_id: int, store_id: int) -> dict[date, Decimal]:
    rows = db.execute(select(PlatformUnit.month, func.sum(PlatformUnit.revenue_cny))
                      .where(PlatformUnit.version_id == version_id, PlatformUnit.store_id == store_id)
                      .group_by(PlatformUnit.month)).all()
    return {m: v for m, v in rows}


def refresh_over_limit(db: Session, version_id: int, store: Store) -> None:
    """按店铺月人民币收入重判超放量；不再超出的月份清除理由（PAGE-F04-02 §7.2）。只在可编辑版本上调用。"""
    start = new_channel_start(db, store)
    limits = ramp_limits(ramp_rule(db))
    revenue = _store_month_revenue(db, version_id, store.id)
    for month, rev in revenue.items():
        over = exceeds_ramp_limit(rev or ZERO, month, start, limits)
        values = {"over_limit": over} if over else {"over_limit": False, "over_limit_reason": None}
        db.execute(update(PlatformUnit).where(PlatformUnit.version_id == version_id,
                                              PlatformUnit.store_id == store.id, PlatformUnit.month == month,
                                              PlatformUnit.over_limit.is_not(over))
                   .values(**values).execution_options(synchronize_session=False))


def over_limit_months(db: Session, version_id: int, store: Store) -> list[dict]:
    start = new_channel_start(db, store)
    limits = ramp_limits(ramp_rule(db)) or []
    rows = db.execute(select(PlatformUnit.month, func.sum(PlatformUnit.revenue_cny),
                             func.max(PlatformUnit.over_limit_reason))
                      .where(PlatformUnit.version_id == version_id, PlatformUnit.store_id == store.id,
                             PlatformUnit.over_limit.is_(True))
                      .group_by(PlatformUnit.month).order_by(PlatformUnit.month)).all()
    out = []
    for month, rev, reason in rows:
        k = new_channel_month_index(month, start)
        limit = limits[k - 1] if k and k <= len(limits) else None
        out.append({"month": fmt_month(month), "k": k, "limit": str(limit) if limit is not None else None,
                    "limit_display": fmt_cny(limit), "revenue_display": fmt_cny(rev), "reason": reason})
    return out


# ---------------------------------------------------------------------------
# 表格：平台渠道（API-F04-02，PAGE-F04-02）
# ---------------------------------------------------------------------------
def _fx_info(db: Session, year: int, currency: str) -> dict:
    rate = db.scalar(select(FxBudget.rate).where(FxBudget.year == year, FxBudget.currency == currency))
    return {"currency": currency, "year": year, "rate": str(rate) if rate is not None else None,
            "missing": rate is None}


def platform_grid(db: Session, ctx: ScopeContext, unit: FillingUnit) -> dict:
    version = _version(db, unit)
    store = db.get(Store, unit.scope_object_id)
    site = db.get(Site, store.site_id)
    channel = db.get(Channel, store.channel_id)
    cur = store.currency
    writable = _can_write(db, ctx, unit)
    if writable:
        refresh_over_limit(db, version.id, store)
    cats = dict(db.execute(select(Category.id, Category.name)).all())
    cat_sort = {cid: i for i, cid in enumerate(db.scalars(select(Category.id).order_by(Category.sort, Category.id)))}
    units = list(db.scalars(select(PlatformUnit).where(PlatformUnit.version_id == version.id,
                                                       PlatformUnit.store_id == store.id)))
    groups: dict[tuple[int, bool], list[PlatformUnit]] = defaultdict(list)
    for u in units:
        groups[(u.category_id, u.is_new)].append(u)
    rows: list[dict] = []
    for (cat_id, is_new) in sorted(groups, key=lambda g: (cat_sort.get(g[0], 0), not g[1])):
        members = sorted(groups[(cat_id, is_new)], key=lambda u: u.month)
        prefix = f"pu:{store.id}:{site.id}:{cat_id}:{1 if is_new else 0}"
        dims = {"category": cats[cat_id], "is_new": "新品" if is_new else "老品"}
        for item, (label, kind) in PLATFORM_INPUTS.items():
            values = {fmt_month(u.month): fmt_input(kind, getattr(u, item)) for u in members}
            rows.append({"key": f"{prefix}|{item}", "kind": "input",
                         "values": {**dims, "item": label.format(cur=cur), **values, "year_total": None}})
        totals = {f: _sum(getattr(u, f) for u in members)
                  for f in ("sales_qty", "revenue_local", "revenue_cny", "gross_profit")}
        totals["gross_margin"] = margin_of(totals["gross_profit"], totals["revenue_cny"])
        for item, label in PLATFORM_CALCS.items():
            values = {fmt_month(u.month): fmt_calc(item, getattr(u, item)) for u in members}
            rows.append({"key": f"{prefix}|{item}", "kind": "calc",
                         "values": {**dims, "item": label.format(cur=cur), **values,
                                    "year_total": fmt_calc(item, totals[item])}})
    # 店铺合计：SQL SUM 聚合未舍入值
    agg = db.execute(select(PlatformUnit.month, func.sum(PlatformUnit.sales_qty), func.sum(PlatformUnit.revenue_local),
                            func.sum(PlatformUnit.revenue_cny), func.sum(PlatformUnit.gross_profit))
                     .where(PlatformUnit.version_id == version.id, PlatformUnit.store_id == store.id)
                     .group_by(PlatformUnit.month)).all()
    by_month = {m: dict(sales_qty=q, revenue_local=rl, revenue_cny=rc, gross_profit=gp) for m, q, rl, rc, gp in agg}
    year_tot = {f: _sum(v[f] for v in by_month.values())
                for f in ("sales_qty", "revenue_local", "revenue_cny", "gross_profit")}
    year_tot["gross_margin"] = margin_of(year_tot["gross_profit"], year_tot["revenue_cny"])
    for item, label in PLATFORM_CALCS.items():
        values = {}
        for m, v in by_month.items():
            values[fmt_month(m)] = fmt_calc(item, margin_of(v["gross_profit"], v["revenue_cny"])
                                            if item == "gross_margin" else v[item])
        rows.append({"key": f"total|{item}", "kind": "total",
                     "values": {"category": "店铺合计", "is_new": "", "item": label.format(cur=cur), **values,
                                "year_total": fmt_calc(item, year_tot[item])}})
    over = over_limit_months(db, version.id, store)
    flags = {f"total|revenue_cny::{o['month']}":
             f"{OVER_LIMIT_LABEL}（起始后第 {o['k']} 个月上限 {o['limit_display']}）" for o in over}
    start = new_channel_start(db, store)
    rv = ramp_rule(db)
    columns = [_dim("category", "品类", 90), _dim("is_new", "新/老品", 72), _dim("item", "项目", 150),
               *_month_cols(version.year, writable)]
    return {
        "unit": T.unit_info(db, ctx, unit),
        "columns": columns,
        "rows": rows,
        "flags": flags,
        "over_limit": over,
        "store": {"id": store.id, "name": store.name, "channel": channel.name, "site": site.name, "currency": cur,
                  "is_new_channel": start is not None, "new_channel_start": fmt_month(start) if start else None},
        "ramp_rule": {"name": "新渠道月收入放量上限", "version_no": rv.version_no,
                      "limits": [str(x) for x in ramp_limits(rv)]} if rv and start else None,
        "fx": _fx_info(db, version.year, cur),
        "kpis": {
            "revenue_local": str(year_tot["revenue_local"] or ZERO),
            "revenue_local_display": fmt_foreign(year_tot["revenue_local"] or ZERO, cur),
            "revenue_cny": str(year_tot["revenue_cny"] or ZERO),
            "revenue_cny_display": fmt_cny(year_tot["revenue_cny"] or ZERO),
            "gross_profit_display": fmt_cny(year_tot["gross_profit"] or ZERO),
            "gross_margin_display": fmt_percent(year_tot["gross_margin"], 1),
        },
        "refs": _platform_refs(db, version.id, store, cats),
    }


def _platform_refs(db: Session, version_id: int, store: Store, cats: dict[int, str]) -> dict:
    """毛利计算引用的比例（只读，DEC-F04-05）：取各比例全年的取值（各月不同时给出区间）。"""
    ccr = db.execute(select(ChannelCategoryRatio).where(ChannelCategoryRatio.version_id == version_id,
                                                        ChannelCategoryRatio.channel_id == store.channel_id,
                                                        ChannelCategoryRatio.site_id == store.site_id)).scalars()
    by_cat: dict[int, dict[str, set]] = defaultdict(lambda: defaultdict(set))
    for r in ccr:
        for f in (*PURCHASE, *LOGISTICS):
            by_cat[r.category_id][f].add(getattr(r, f))
    csr = list(db.scalars(select(ChannelSiteRatio).where(ChannelSiteRatio.version_id == version_id,
                                                        ChannelSiteRatio.channel_id == store.channel_id,
                                                        ChannelSiteRatio.site_id == store.site_id)))

    def span(vals: set) -> str:
        if not vals:
            return "未填"
        lo, hi = min(vals), max(vals)
        return fmt_percent(lo, 1) if lo == hi else f"{fmt_percent(lo, 1)}～{fmt_percent(hi, 1)}"

    return {
        "categories": [{"category": cats[cid], **{f: span(v[f]) for f in (*PURCHASE, *LOGISTICS)}}
                       for cid, v in sorted(by_cat.items())],
        "site": {f: span({getattr(r, f) for r in csr}) for f in SITE_RATES},
    }


# ---------------------------------------------------------------------------
# 表格：B 端（PAGE-F04-03）
# ---------------------------------------------------------------------------
def _b2b_customers(db: Session, unit: FillingUnit) -> list[BCustomer]:
    ids = set(db.scalars(select(B2bUnit.customer_id).where(B2bUnit.version_id == unit.version_id)))
    return list(db.scalars(select(BCustomer).where(BCustomer.department_id == unit.scope_object_id,
                                                   BCustomer.id.in_(ids)).order_by(BCustomer.id)))


def b2b_grid(db: Session, ctx: ScopeContext, unit: FillingUnit) -> dict:
    version = _version(db, unit)
    writable = _can_write(db, ctx, unit)
    customers = [c for c in _b2b_customers(db, unit) if ctx.all_company or c.id in ctx.customer_ids
                 or unit.reviewer_id == ctx.user_id or ctx.has_role(Role.DEPT_HEAD)]
    rows: list[dict] = []
    for c in customers:
        members = list(db.scalars(select(B2bUnit).where(B2bUnit.version_id == version.id,
                                                        B2bUnit.customer_id == c.id).order_by(B2bUnit.month)))
        editable_cust = writable and (ctx.all_company or c.id in ctx.customer_ids)
        dims = {"customer": c.name}
        for item, (label, kind) in B2B_INPUTS.items():
            values = {fmt_month(u.month): fmt_input(kind, getattr(u, item)) for u in members}
            rows.append({"key": f"b2b:{c.id}|{item}", "kind": "input",
                         "readonly_fields": [] if editable_cust else list(values),
                         "values": {**dims, "item": label.format(cur=c.currency), **values,
                                    "year_total": fmt_amount(_sum(u.sales_local for u in members))
                                    if item == "sales_local" else None}})
        rev, gp = _sum(u.revenue_cny for u in members), _sum(u.gross_profit for u in members)
        tot = {"revenue_cny": rev, "gross_profit": gp, "gross_margin": margin_of(gp, rev)}
        for item, label in B2B_CALCS.items():
            values = {fmt_month(u.month): fmt_calc(item, getattr(u, item)) for u in members}
            rows.append({"key": f"b2b:{c.id}|{item}", "kind": "calc",
                         "values": {**dims, "item": label, **values, "year_total": fmt_calc(item, tot[item])}})
    cust_ids = [c.id for c in customers]
    currencies = sorted({c.currency for c in customers})
    year_local: dict[str, Decimal | None] = {}
    if cust_ids:
        for cur in currencies:
            agg = db.execute(select(B2bUnit.month, func.sum(B2bUnit.sales_local)).join(
                BCustomer, BCustomer.id == B2bUnit.customer_id)
                .where(B2bUnit.version_id == version.id, B2bUnit.customer_id.in_(cust_ids), BCustomer.currency == cur)
                .group_by(B2bUnit.month)).all()
            year_local[cur] = _sum(v for _, v in agg)
            rows.append({"key": f"total:{cur}|sales_local", "kind": "total",
                         "values": {"customer": "合计", "item": f"销售额（{cur}）",
                                    **{fmt_month(m): fmt_amount(v) for m, v in agg},
                                    "year_total": fmt_amount(year_local[cur])}})
        agg = db.execute(select(B2bUnit.month, func.sum(B2bUnit.revenue_cny), func.sum(B2bUnit.gross_profit))
                         .where(B2bUnit.version_id == version.id, B2bUnit.customer_id.in_(cust_ids))
                         .group_by(B2bUnit.month)).all()
    else:
        agg = []
    rev_y, gp_y = _sum(r for _, r, _ in agg), _sum(g for _, _, g in agg)
    for item, label in B2B_CALCS.items():
        values = {}
        for m, r, g in agg:
            values[fmt_month(m)] = fmt_calc(item, {"revenue_cny": r, "gross_profit": g,
                                                   "gross_margin": margin_of(g, r)}[item])
        yt = {"revenue_cny": rev_y, "gross_profit": gp_y, "gross_margin": margin_of(gp_y, rev_y)}[item]
        rows.append({"key": f"total|{item}", "kind": "total",
                     "values": {"customer": "合计", "item": label, **values, "year_total": fmt_calc(item, yt)}})
    refs = []
    for c in customers:
        members = list(db.scalars(select(B2bUnit).where(B2bUnit.version_id == version.id,
                                                        B2bUnit.customer_id == c.id)))
        refs.append({"customer": c.name, **{f: _span(getattr(u, f) for u in members) for f in (*PURCHASE, *LOGISTICS)}})
    dept = db.get(Department, unit.scope_object_id)
    return {
        "unit": T.unit_info(db, ctx, unit),
        "columns": [_dim("customer", "B 端客户", 110), _dim("item", "项目", 150), *_month_cols(version.year, writable)],
        "rows": rows,
        "flags": {},
        "department": dept.name if dept else None,
        "customers": [{"id": c.id, "name": c.name, "currency": c.currency} for c in customers],
        "fx": [_fx_info(db, version.year, cur) for cur in currencies],
        "kpis": {
            "sales_local_display": "、".join(fmt_foreign(year_local.get(cur) or ZERO, cur) for cur in currencies) or "—",
            "revenue_cny_display": fmt_cny(rev_y or ZERO),
            "gross_profit_display": fmt_cny(gp_y or ZERO),
            "gross_margin_display": fmt_percent(margin_of(gp_y, rev_y), 1),
        },
        "refs": refs,
    }


def _span(values) -> str:
    vals = {v for v in values if v is not None}
    if not vals:
        return "未填"
    lo, hi = min(vals), max(vals)
    return fmt_percent(lo, 1) if lo == hi else f"{fmt_percent(lo, 1)}～{fmt_percent(hi, 1)}"


# ---------------------------------------------------------------------------
# 表格：成本比例 / 物流比例（PAGE-F04-04、PAGE-F04-05）
# ---------------------------------------------------------------------------
def _names(db: Session, model) -> dict[int, str]:
    return dict(db.execute(select(model.id, model.name)).all())


def _ratio_rows(db: Session, version: BudgetVersion, prefix: str, items: dict) -> list[dict]:
    rows: list[dict] = []
    if prefix == "ccr":
        channels, sites, cats = _names(db, Channel), _names(db, Site), _names(db, Category)
        ch_sort = dict(db.execute(select(Channel.id, Channel.sort)).all())
        cat_sort = dict(db.execute(select(Category.id, Category.sort)).all())
        data = list(db.scalars(select(ChannelCategoryRatio).where(ChannelCategoryRatio.version_id == version.id)))
        groups: dict[tuple, list] = defaultdict(list)
        for r in data:
            groups[(r.channel_id, r.site_id, r.category_id)].append(r)
        for key in sorted(groups, key=lambda k: (ch_sort.get(k[0], 0), k[0], k[1], cat_sort.get(k[2], 0), k[2])):
            members = sorted(groups[key], key=lambda r: r.month)
            dims = {"channel": channels[key[0]], "site": sites[key[1]], "category": cats[key[2]],
                    "channel_id": str(key[0])}
            for item, (label, kind) in items.items():
                rows.append({"key": f"ccr:{key[0]}:{key[1]}:{key[2]}|{item}", "kind": "input",
                             "values": {**dims, "item": label,
                                        **{fmt_month(r.month): fmt_input(kind, getattr(r, item)) for r in members}}})
    elif prefix == "b2bc":
        customers = {c.id: c for c in db.scalars(select(BCustomer))}
        depts = _names(db, Department)
        data = list(db.scalars(select(B2bUnit).where(B2bUnit.version_id == version.id)))
        groups = defaultdict(list)
        for r in data:
            groups[r.customer_id].append(r)
        for cid in sorted(groups):
            members = sorted(groups[cid], key=lambda r: r.month)
            c = customers[cid]
            dims = {"customer": c.name, "department": depts.get(c.department_id, "")}
            for item, (label, kind) in items.items():
                rows.append({"key": f"b2bc:{cid}|{item}", "kind": "input",
                             "values": {**dims, "item": label,
                                        **{fmt_month(r.month): fmt_input(kind, getattr(r, item)) for r in members}}})
    elif prefix == "csr":
        channels, sites = _names(db, Channel), _names(db, Site)
        ch_sort = dict(db.execute(select(Channel.id, Channel.sort)).all())
        data = list(db.scalars(select(ChannelSiteRatio).where(ChannelSiteRatio.version_id == version.id)))
        groups = defaultdict(list)
        for r in data:
            groups[(r.channel_id, r.site_id)].append(r)
        for key in sorted(groups, key=lambda k: (ch_sort.get(k[0], 0), k[0], k[1])):
            members = sorted(groups[key], key=lambda r: r.month)
            dims = {"channel": channels[key[0]], "site": sites[key[1]], "channel_id": str(key[0])}
            for item, (label, kind) in items.items():
                rows.append({"key": f"csr:{key[0]}:{key[1]}|{item}", "kind": "input",
                             "values": {**dims, "item": label,
                                        **{fmt_month(r.month): fmt_input(kind, getattr(r, item)) for r in members}}})
    return rows


def _other_pl_rows(db: Session, version: BudgetVersion) -> list[dict]:
    data = list(db.scalars(select(CompanyOtherPl).where(CompanyOtherPl.version_id == version.id)))
    rows = []
    for item, label in OtherPlItem.LABELS.items():
        members = sorted((r for r in data if r.item == item), key=lambda r: r.month)
        rows.append({"key": f"opl|{item}", "kind": "input",
                     "values": {"item": label, **{fmt_month(r.month): fmt_input("signed2", r.amount) for r in members},
                                "year_total": fmt_amount(_sum(r.amount for r in members) or ZERO)}})
    agg = db.execute(select(CompanyOtherPl.month, func.sum(CompanyOtherPl.amount))
                     .where(CompanyOtherPl.version_id == version.id).group_by(CompanyOtherPl.month)).all()
    rows.append({"key": "total|other_pl", "kind": "total",
                 "values": {"item": "其他损益合计", **{fmt_month(m): fmt_amount(v) for m, v in agg},
                            "year_total": fmt_amount(_sum(v for _, v in agg) or ZERO)}})
    return rows


def ratio_grid(db: Session, ctx: ScopeContext, unit: FillingUnit, tab: str) -> dict:
    version = _version(db, unit)
    writable = _can_write(db, ctx, unit)
    logistics = unit.unit_type == UnitType.LOGISTICS_RATIO
    if tab not in TABS[unit.unit_type]:
        raise ValidationFailed("标签不存在", code="invalid_tab")
    items = LOGISTICS if logistics else PURCHASE
    if tab == "platform":
        dims = [_dim("channel", "渠道", 100), _dim("site", "站点", 70), _dim("category", "品类", 80)]
        if logistics:
            dims.append(_dim("item", "比例", 100))
        rows = _ratio_rows(db, version, "ccr", items)
        columns = [*dims, *_month_cols(version.year, writable, total_label=None)]
    elif tab == "b2b":
        dims = [_dim("customer", "B 端客户", 110), _dim("department", "负责部门", 100)]
        if logistics:
            dims.append(_dim("item", "比例", 100))
        rows = _ratio_rows(db, version, "b2bc", items)
        columns = [*dims, *_month_cols(version.year, writable, total_label=None)]
    elif tab == "site":
        rows = _ratio_rows(db, version, "csr", SITE_RATES)
        columns = [_dim("channel", "渠道", 100), _dim("site", "站点", 70), _dim("item", "项目", 130),
                   *_month_cols(version.year, writable, total_label=None)]
    else:  # other：公司层其他损益（AC-F06-06）
        rows = _other_pl_rows(db, version)
        columns = [_dim("item", "项目", 140), *_month_cols(version.year, writable, total_label="全年合计")]
    return {"unit": T.unit_info(db, ctx, unit), "tab": tab, "tabs": list(TABS[unit.unit_type]),
            "columns": columns, "rows": rows, "flags": {}}


def adjustments_module():
    """G3 的目标调整服务（安全导入：缺失时调整模式不可用，编制填报不受影响）。"""
    try:
        return importlib.import_module("app.services.adjustments")
    except ImportError:
        return None


def versions_module():
    """G3 的版本服务（安全导入）：open_months / plan_units / submit_rolling_unit。"""
    try:
        return importlib.import_module("app.services.versions")
    except ImportError:
        return None


def locked_month_keys(db: Session, unit: FillingUnit, year: int) -> set[str]:
    """调整稿：已导入金蝶实际的月份不可修改（AC-F08-02）；滚动预测稿：只开放本季度及以后月份（AC-F08-05）；
    编制稿全部月份可编辑。以 G3 versions.open_months 为准。"""
    version = db.get(BudgetVersion, unit.version_id)
    if version.type not in (VersionType.ADJUSTMENT_DRAFT, VersionType.ROLLING_DRAFT):
        return set()
    vmod = versions_module()
    if vmod is not None:
        months = vmod.open_months(db, unit.version_id)
        if months is None:
            return set()
        return {fmt_month(m) for m in year_months(year)} - {fmt_month(m) for m in months}
    mod = adjustments_module() if unit.adjustment_id else None
    if mod is None:
        return set()
    return {fmt_month(m) for m in year_months(year)} - {fmt_month(m) for m in mod.adjustable_months(db, year)}


def month_label_suffix(db: Session, unit: FillingUnit) -> str:
    return "（实际）" if db.get(BudgetVersion, unit.version_id).type == VersionType.ADJUSTMENT_DRAFT else "（未开放）"


def build_grid(db: Session, ctx: ScopeContext, unit: FillingUnit, tab: str | None = None) -> dict:
    T.ensure_view(db, ctx, unit)
    if unit.unit_type == UnitType.LOGISTICS_RATIO and not ctx.has_role(Role.LOGISTICS) and not ctx.all_company \
            and unit.reviewer_id != ctx.user_id:
        raise Forbidden()
    if unit.unit_type == UnitType.PLATFORM:
        grid = platform_grid(db, ctx, unit)
    elif unit.unit_type == UnitType.B2B:
        grid = b2b_grid(db, ctx, unit)
    elif unit.unit_type == UnitType.DEPT_EXPENSE:
        from app.services.dept_expense import dept_grid
        grid = dept_grid(db, ctx, unit)
    else:
        grid = ratio_grid(db, ctx, unit, tab or TABS[unit.unit_type][0])
    locked = locked_month_keys(db, unit, _version(db, unit).year)
    grid["actual_months"] = sorted(locked)
    if locked:
        suffix = month_label_suffix(db, unit)
        for col in grid["columns"]:
            if col["key"] in locked:
                col["label"] += suffix
        for row in grid["rows"]:
            if row.get("kind") == "input":
                row["readonly_fields"] = sorted(set(row.get("readonly_fields") or []) | locked)
    return grid


# ---------------------------------------------------------------------------
# 单元格保存（API-F04-03，SEQ-F04-01，DEC-007）
# ---------------------------------------------------------------------------
@dataclass
class Target:
    model: type
    row: Any
    attr: str
    kind: str


def _parse_key(row_key: str) -> tuple[str, list[str], str]:
    head, _, item = row_key.partition("|")
    parts = head.split(":")
    return parts[0], parts[1:], item


def _to_int(parts: list[str]) -> list[int]:
    try:
        return [int(p) for p in parts]
    except ValueError:
        raise NotFound("单元格不存在") from None


def _locate(db: Session, unit: FillingUnit, ctx: ScopeContext, prefix: str, ids: list[int], month: date,
            item: str, *, lock: bool) -> Any:
    """按行 key 找到预算数据行，并校验该行属于本单位、在本人范围内。"""
    v = unit.version_id
    if prefix == "pu":
        if len(ids) != 4 or ids[0] != unit.scope_object_id:
            raise Forbidden()
        q = select(PlatformUnit).where(PlatformUnit.version_id == v, PlatformUnit.store_id == ids[0],
                                       PlatformUnit.site_id == ids[1], PlatformUnit.category_id == ids[2],
                                       PlatformUnit.is_new.is_(bool(ids[3])), PlatformUnit.month == month)
    elif prefix in ("b2b", "b2bc"):
        if len(ids) != 1:
            raise NotFound("单元格不存在")
        if prefix == "b2b":
            cust = db.get(BCustomer, ids[0])
            if cust is None or cust.department_id != unit.scope_object_id:
                raise Forbidden()
            ctx.ensure_customer(cust.id)
        q = select(B2bUnit).where(B2bUnit.version_id == v, B2bUnit.customer_id == ids[0], B2bUnit.month == month)
    elif prefix == "ccr":
        if len(ids) != 3:
            raise NotFound("单元格不存在")
        q = select(ChannelCategoryRatio).where(ChannelCategoryRatio.version_id == v,
                                               ChannelCategoryRatio.channel_id == ids[0],
                                               ChannelCategoryRatio.site_id == ids[1],
                                               ChannelCategoryRatio.category_id == ids[2],
                                               ChannelCategoryRatio.month == month)
    elif prefix == "csr":
        if len(ids) != 2:
            raise NotFound("单元格不存在")
        q = select(ChannelSiteRatio).where(ChannelSiteRatio.version_id == v, ChannelSiteRatio.channel_id == ids[0],
                                           ChannelSiteRatio.site_id == ids[1], ChannelSiteRatio.month == month)
    elif prefix == "opl":
        q = select(CompanyOtherPl).where(CompanyOtherPl.version_id == v, CompanyOtherPl.item == item,
                                         CompanyOtherPl.month == month)
    elif prefix == "de":
        if len(ids) != 1:
            raise NotFound("单元格不存在")
        q = select(DeptExpense).where(DeptExpense.version_id == v, DeptExpense.department_id == unit.scope_object_id,
                                      DeptExpense.account_id == ids[0], DeptExpense.month == month)
    else:
        raise NotFound("单元格不存在")
    if lock:
        q = q.with_for_update().execution_options(populate_existing=True)
    row = db.scalar(q)
    if row is None:
        raise NotFound("单元格不存在")
    return row


def _stored_attr(prefix: str, item: str) -> str:
    return "amount" if prefix == "opl" else item


def _current_display(kind: str, value) -> str:
    return fmt_input(kind, value)


def _same(kind: str, base: Any, current) -> bool:
    """基准值与当前值是否相同（按数值比较，容忍格式差异）。基准值缺省按 0。"""
    try:
        b = parse_cell(kind, base)
    except CellError:
        return False
    cur = current if current is not None else (0 if kind == "int" else ZERO)
    return Decimal(b) == Decimal(cur)


def _resolve(db: Session, unit_or_id: FillingUnit | int, *, lock: bool = False) -> FillingUnit:
    """填报单位：库中单位按 ID 取（可加行锁）；滚动预测单位为按 G3 plan_units 推导的临时对象，原样返回。"""
    if isinstance(unit_or_id, FillingUnit):
        return unit_or_id
    return T.get_unit(db, unit_or_id, lock=lock)


def _audit_id(unit: FillingUnit) -> str | int:
    return unit.id if unit.id else f"rolling:{unit.version_id}:{getattr(unit, 'rolling_key', '')}"


def save_cells(db: Session, ctx: ScopeContext, actor: CurrentUser, unit_id: FillingUnit | int, changes: list[dict],
               *, check_base: bool = True, request=None, source: str = "edit") -> dict:
    unit = _resolve(db, unit_id, lock=True)
    T.ensure_writable(db, ctx, unit)
    version = _version(db, unit)
    allowed = WRITABLE[unit.unit_type]
    now = clock.now()
    errors: list[dict] = []
    ignored: list[dict] = []
    saved: list[dict] = []
    pu_ids: set[int] = set()
    b2b_ids: set[int] = set()
    ccr_keys: set[tuple] = set()
    csr_keys: set[tuple] = set()
    touched_store = False
    locked = locked_month_keys(db, unit, version.year)
    adj_mod = adjustments_module() if unit.adjustment_id else None
    for ch in changes:
        row_key, field = str(ch.get("row_key") or ""), str(ch.get("field") or "")
        prefix, parts, item = _parse_key(row_key)
        items = allowed.get(prefix)
        if items is None or item not in items:
            if item in ALL_INPUT_ITEMS and item not in (items or {}):
                raise Forbidden()  # 其他角色的输入项（如物流物控写采购成本占比，AC-F01-09）
            ignored.append({"row_key": row_key, "field": field})  # 派生字段 / 未知字段：伪造即忽略（RULE-01）
            continue
        try:
            month = parse_month(field)
        except ValueError:
            ignored.append({"row_key": row_key, "field": field})
            continue
        if month.year != version.year:
            ignored.append({"row_key": row_key, "field": field})
            continue
        if unit.adjustment_id is not None and adj_mod is not None:
            try:
                adj_mod.ensure_month_adjustable(db, unit, month)  # AC-F08-02：「该月已导入实际，不能调整」
            except ValidationFailed as e:
                errors.append({"row_key": row_key, "field": field, "code": "month_locked", "message": e.message})
                continue
        if field in locked:
            errors.append({"row_key": row_key, "field": field, "code": "month_locked",
                           "message": "已导入金蝶实际，不能调整" if unit.adjustment_id
                           else "滚动预测只开放本季度及以后的月份"})
            continue
        label, kind = items[item]
        try:
            value = parse_cell(kind, ch.get("value"))
        except CellError as e:
            errors.append({"row_key": row_key, "field": field, "code": "out_of_range", "message": e.grid_message,
                           "range": RANGES[kind]})
            continue
        ids = _to_int(parts) if prefix != "opl" else []
        row = _locate(db, unit, ctx, prefix, ids, month, item, lock=True)
        attr = _stored_attr(prefix, item)
        current = getattr(row, attr)
        if check_base and not _same(kind, ch.get("base"), current):
            who = T.user_name(db, row.updated_by) or "其他人"
            errors.append({"row_key": row_key, "field": field, "code": "conflict",
                           "message": f"该数据已被 {who} 于 {T.fmt_dt(row.updated_at)} 修改，请刷新后再改",
                           "current": _current_display(kind, current)})
            continue
        if Decimal(current if current is not None else 0) == Decimal(value):
            saved.append({"row_key": row_key, "field": field})
            continue
        setattr(row, attr, value)
        row.updated_by = actor.user.id
        row.updated_at = now
        saved.append({"row_key": row_key, "field": field, "before": current, "after": value})
        if prefix == "pu":
            pu_ids.add(row.id)
            touched_store = True
        elif prefix in ("b2b", "b2bc"):
            b2b_ids.add(row.id)
        elif prefix == "ccr":
            ccr_keys.add((row.month, row.channel_id, row.site_id, row.category_id))
        elif prefix == "csr":
            csr_keys.add((row.month, row.channel_id, row.site_id))
    db.flush()
    # 增量重算：只算受影响的单元（DEC-006）
    if pu_ids:
        recalc_platform_units(db, version.id, unit_ids=sorted(pu_ids))
    for m, c, s, cat in sorted(ccr_keys):
        recalc_platform_units(db, version.id, month=m, channel_id=c, site_id=s, category_id=cat)
    for m, c, s in sorted(csr_keys):
        recalc_platform_units(db, version.id, month=m, channel_id=c, site_id=s)
    if b2b_ids:
        recalc_b2b_units(db, version.id, unit_ids=sorted(b2b_ids))
    if touched_store:
        refresh_over_limit(db, version.id, db.get(Store, unit.scope_object_id))
    changed = [s for s in saved if "after" in s]
    if changed:
        T.mark_filling(unit)
        write_audit(db, action=f"cells.{source}", object_type="filling_unit", object_id=_audit_id(unit), actor=actor,
                    request=request, before=[{"row_key": s["row_key"], "field": s["field"], "value": s["before"]}
                                             for s in changed],
                    after=[{"row_key": s["row_key"], "field": s["field"], "value": s["after"]} for s in changed])
    db.flush()
    return {"saved": len(saved), "errors": errors, "ignored": ignored}


def fill_year(db: Session, ctx: ScopeContext, actor: CurrentUser, unit_id: FillingUnit | int, row_key: str,
              value: Any, request=None) -> dict:
    """AC-F04-06：把一个比率写入该维度组合 12 个月中尚未提交的月份（调整 / 滚动的锁定月份由版本服务限制）。"""
    unit = _resolve(db, unit_id)
    T.ensure_writable(db, ctx, unit)
    prefix, _, item = _parse_key(row_key)
    items = WRITABLE[unit.unit_type].get(prefix) or {}
    if item not in items or items[item][1] != "rate":
        raise ValidationFailed("只能对比率行填充全年", code="not_ratio")
    try:
        parse_cell("rate", value)
    except CellError as e:
        raise ValidationFailed(e.grid_message, code="out_of_range") from None
    version = _version(db, unit)
    changes = [{"row_key": row_key, "field": fmt_month(m), "value": value} for m in year_months(version.year)]
    return save_cells(db, ctx, actor, unit, changes, check_base=False, request=request, source="fill_year")


# ---------------------------------------------------------------------------
# 新渠道超出理由（API-F04-08）与提交（API-F04-07）
# ---------------------------------------------------------------------------
def set_over_limit_reasons(db: Session, ctx: ScopeContext, actor: CurrentUser, unit_id: int,
                           reasons: dict[str, str | None], request=None) -> list[dict]:
    unit = T.get_unit(db, unit_id, lock=True)
    T.ensure_writable(db, ctx, unit)
    if unit.unit_type != UnitType.PLATFORM:
        raise ValidationFailed("只有平台渠道单位需要填写超出理由", code="not_platform")
    store = db.get(Store, unit.scope_object_id)
    refresh_over_limit(db, unit.version_id, store)
    over = {o["month"] for o in over_limit_months(db, unit.version_id, store)}
    for month_key, text in reasons.items():
        if month_key not in over:
            raise ValidationFailed(f"{month_key} 未超出新渠道放量规则，无需填写理由", code="not_over_limit")
        month = parse_month(month_key)
        db.execute(update(PlatformUnit).where(PlatformUnit.version_id == unit.version_id,
                                              PlatformUnit.store_id == store.id, PlatformUnit.month == month)
                   .values(over_limit_reason=(text or "").strip() or None)
                   .execution_options(synchronize_session=False))
    T.mark_filling(unit)
    write_audit(db, action="unit.over_limit_reasons", object_type="filling_unit", object_id=unit.id, actor=actor,
                request=request, after=reasons)
    db.flush()
    db.expire_all()
    return over_limit_months(db, unit.version_id, store)


def _unit_currencies(db: Session, unit: FillingUnit) -> list[str]:
    if unit.unit_type == UnitType.PLATFORM:
        return [db.scalar(select(Store.currency).where(Store.id == unit.scope_object_id))]
    if unit.unit_type == UnitType.B2B:
        return sorted({c.currency for c in _b2b_customers(db, unit)})
    return []


def submit_unit(db: Session, ctx: ScopeContext, actor: CurrentUser, settings, unit_id: int, request=None,
                adjust_reason: str | None = None) -> dict:
    """提交：汇率齐全（AC-F02-07）、超放量已写理由（AC-F13-02）→ 已提交并锁定（AC-F04-09）→ 通知审核人（AC-F07-01）。
    调整单位交给目标调整服务提交（调整原因必填，AC-F08-03）。"""
    unit = T.get_unit(db, unit_id, lock=True)
    T.ensure_writable(db, ctx, unit)
    if unit.adjustment_id is not None:
        mod = adjustments_module()
        if mod is None:
            raise Conflict("目标调整暂不可用", code="adjustment_unavailable")
        mod.submit_unit(db, actor, unit.id, adjust_reason)
        db.refresh(unit)
        return T.unit_info(db, ctx, unit)
    if unit.status == UnitStatus.NOT_STARTED:
        raise ValidationFailed("本单位尚未填写任何数据，不能提交", code="not_started")
    version = _version(db, unit)
    have = set(db.scalars(select(FxBudget.currency).where(FxBudget.year == version.year)))
    missing = [c for c in _unit_currencies(db, unit) if c not in have]
    if missing:
        raise ValidationFailed(f"{version.year} 年度缺少以下币种的预算汇率：{'、'.join(missing)}。"
                               "请联系财务BP 在「汇率」中录入后再提交", code="missing_fx",
                               details={"currencies": missing})
    if unit.unit_type == UnitType.PLATFORM:
        store = db.get(Store, unit.scope_object_id)
        refresh_over_limit(db, version.id, store)
        lacking = [o["month"] for o in over_limit_months(db, version.id, store) if not o["reason"]]
        if lacking:
            raise ValidationFailed(f"以下月份超出新渠道放量规则，请填写超出理由后再提交：{'、'.join(lacking)}",
                                   code="over_limit_reason_required", details={"months": lacking})
    before = unit.status
    res = db.execute(update(FillingUnit).where(FillingUnit.id == unit.id,
                                               FillingUnit.status.in_(UnitStatus.EDITABLE))
                     .values(status=UnitStatus.SUBMITTED, submitted_at=clock.now(), updated_at=clock.now())
                     .execution_options(synchronize_session=False))
    if res.rowcount == 0:
        raise Conflict("状态已变化，请刷新", code="status_changed")
    db.refresh(unit)
    write_audit(db, action="unit.submit", object_type="filling_unit", object_id=unit.id, actor=actor,
                request=request, before={"status": before}, after={"status": unit.status})
    task = T.task_of(db, unit)
    T.notify(db, biz_type="unit_submitted", biz_id=unit.id, user_ids=[unit.reviewer_id],
             title="【预算管理平台】有填报待你审核", name=task.name if task else "预算填报",
             object_name=unit.name, period=f"{version.year} 年度",
             link=T.app_link(settings, f"/approvals/unit/{unit.id}"))
    db.flush()
    return T.unit_info(db, ctx, unit)


# ---------------------------------------------------------------------------
# 我的填报（API-F04-01）
# ---------------------------------------------------------------------------
def my_units(db: Session, ctx: ScopeContext) -> list[dict]:
    tasks = list(db.scalars(select(BudgetTask).order_by(BudgetTask.year.desc(), BudgetTask.id.desc())))
    latest_year = tasks[0].year if tasks else None
    visible_tasks = [t.id for t in tasks if t.status != "completed" or t.year == latest_year]
    open_adjustments = list(db.scalars(select(TargetAdjustment.id).where(
        TargetAdjustment.status.in_((AdjustmentStatus.IN_PROGRESS, AdjustmentStatus.PENDING_CEO)))))
    if not visible_tasks and not open_adjustments:
        return my_rolling_units(db, ctx)
    units = db.scalars(select(FillingUnit).where(or_(FillingUnit.task_id.in_(visible_tasks),
                                                     FillingUnit.adjustment_id.in_(open_adjustments)))
                       .order_by(FillingUnit.id))
    mine = [u for u in units if T.can_fill(db, ctx, u)]
    mine.sort(key=lambda u: (T.STATUS_SORT.get(u.status, 9), T.UNIT_TYPE_ORDER.index(u.unit_type), u.id))
    out = []
    for u in mine:
        info = T.unit_info(db, ctx, u)
        info["path"] = T.fill_path(u)
        out.append(info)
    return out + my_rolling_units(db, ctx)


# ---------------------------------------------------------------------------
# Excel 模板（API-F04-05 / API-F04-06，AC-F04-07/08）
# ---------------------------------------------------------------------------
@dataclass
class Sheet:
    title: str
    dims: list[str]          # 维度列表头（不含「月份」）
    items: dict              # item -> (label, kind)
    prefix: str


def _sheets(db: Session, unit: FillingUnit) -> list[Sheet]:
    t = unit.unit_type
    if t == UnitType.PLATFORM:
        cur = db.scalar(select(Store.currency).where(Store.id == unit.scope_object_id))
        items = {k: (v[0].format(cur=cur), v[1]) for k, v in PLATFORM_INPUTS.items()}
        return [Sheet("平台渠道驱动", ["店铺", "站点", "品类", "新老品"], items, "pu")]
    if t == UnitType.B2B:
        return [Sheet("B端客户", ["B 端客户"], {k: (v[0].format(cur="本币"), v[1]) for k, v in B2B_INPUTS.items()},
                      "b2b")]
    if t == UnitType.COST_RATIO:
        return [Sheet("采购成本占比·平台", ["渠道", "站点", "品类"], PURCHASE, "ccr"),
                Sheet("采购成本占比·B端客户", ["B 端客户"], PURCHASE, "b2bc"),
                Sheet("库存索赔与其他运营费用率", ["渠道", "站点"], SITE_RATES, "csr"),
                Sheet("公司层其他损益", [], OTHER_PL, "opl")]
    if t == UnitType.LOGISTICS_RATIO:
        return [Sheet("平台渠道", ["渠道", "站点", "品类"], LOGISTICS, "ccr"),
                Sheet("B端客户", ["B 端客户"], LOGISTICS, "b2bc")]
    return [Sheet("部门费用", ["预算科目"], {"amount": ("预算金额（人民币）", "money2")}, "de")]


def _col_header(label: str, kind: str) -> str:
    return f"{label}（%）" if kind == "rate" else label


def _template_rows(db: Session, ctx: ScopeContext, unit: FillingUnit, sheet: Sheet) -> list[tuple[list[str], str, date, Any]]:
    """[(维度值, 行 key 前缀（不含 |item）, 月份, 数据行)]，只含本人范围内的维度组合。"""
    v = unit.version_id
    out = []
    if sheet.prefix == "pu":
        store = db.get(Store, unit.scope_object_id)
        site = db.get(Site, store.site_id)
        cats = _names(db, Category)
        cat_sort = dict(db.execute(select(Category.id, Category.sort)).all())
        rows = db.scalars(select(PlatformUnit).where(PlatformUnit.version_id == v,
                                                     PlatformUnit.store_id == store.id))
        for r in sorted(rows, key=lambda r: (cat_sort[r.category_id], not r.is_new, r.month)):
            out.append(([store.name, site.name, cats[r.category_id], "新品" if r.is_new else "老品"],
                        f"pu:{store.id}:{site.id}:{r.category_id}:{1 if r.is_new else 0}", r.month, r))
    elif sheet.prefix in ("b2b", "b2bc"):
        custs = {c.id: c for c in db.scalars(select(BCustomer))}
        q = select(B2bUnit).where(B2bUnit.version_id == v)
        if sheet.prefix == "b2b":
            q = q.where(B2bUnit.customer_id.in_([c.id for c in _b2b_customers(db, unit)]))
        for r in sorted(db.scalars(q), key=lambda r: (r.customer_id, r.month)):
            if sheet.prefix == "b2b" and not (ctx.all_company or r.customer_id in ctx.customer_ids):
                continue
            out.append(([custs[r.customer_id].name], f"{sheet.prefix}:{r.customer_id}", r.month, r))
    elif sheet.prefix == "ccr":
        ch, si, ca = _names(db, Channel), _names(db, Site), _names(db, Category)
        for r in sorted(db.scalars(select(ChannelCategoryRatio).where(ChannelCategoryRatio.version_id == v)),
                        key=lambda r: (r.channel_id, r.site_id, r.category_id, r.month)):
            out.append(([ch[r.channel_id], si[r.site_id], ca[r.category_id]],
                        f"ccr:{r.channel_id}:{r.site_id}:{r.category_id}", r.month, r))
    elif sheet.prefix == "csr":
        ch, si = _names(db, Channel), _names(db, Site)
        for r in sorted(db.scalars(select(ChannelSiteRatio).where(ChannelSiteRatio.version_id == v)),
                        key=lambda r: (r.channel_id, r.site_id, r.month)):
            out.append(([ch[r.channel_id], si[r.site_id]], f"csr:{r.channel_id}:{r.site_id}", r.month, r))
    elif sheet.prefix == "opl":
        rows = {(r.item, r.month): r for r in db.scalars(select(CompanyOtherPl).where(CompanyOtherPl.version_id == v))}
        months = sorted({m for _, m in rows})
        for m in months:
            out.append(([], "opl", m, {item: rows.get((item, m)) for item in OtherPlItem.LABELS}))
    elif sheet.prefix == "de":
        accs = _names(db, BudgetAccount)
        for r in sorted(db.scalars(select(DeptExpense).where(DeptExpense.version_id == v,
                                                              DeptExpense.department_id == unit.scope_object_id)),
                        key=lambda r: (r.account_id, r.month)):
            out.append(([accs[r.account_id]], f"de:{r.account_id}", r.month, r))
    return out


def _template_value(sheet: Sheet, row: Any, item: str) -> str:
    """模板预填值：与表格输入格同一无损口径（比率为百分数，不带「%」与千分位）。"""
    kind = sheet.items[item][1]
    if sheet.prefix == "opl":
        rec = row.get(item)
        v = rec.amount if rec else ZERO
    else:
        v = getattr(row, item)
    return fmt_input(kind, v).replace(",", "").replace("%", "").replace(MINUS, "-")


def build_template(db: Session, ctx: ScopeContext, unit: FillingUnit) -> tuple[bytes, str]:
    T.ensure_view(db, ctx, unit)
    if not T.can_fill(db, ctx, unit) and not ctx.all_company:
        raise Forbidden()
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in _sheets(db, unit):
        ws = wb.create_sheet(sheet.title)
        ws.append([*sheet.dims, "月份", *(_col_header(lbl, kind) for lbl, kind in sheet.items.values())])
        for dims, _, month, row in _template_rows(db, ctx, unit, sheet):
            ws.append([*dims, fmt_month(month), *(_template_value(sheet, row, item) for item in sheet.items)])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    year = _version(db, unit).year
    return buf.getvalue(), f"{unit.name}_{year}年度预算填报模板.xlsx"


def upload_template(db: Session, ctx: ScopeContext, actor: CurrentUser, unit_id: int, file_name: str,
                    content: bytes, request=None) -> dict:
    """整份校验：任一行不通过则不写入任何数据，列出全部错误的行号、字段、原因（AC-F04-08、DEC-008）。"""
    unit = T.get_unit(db, unit_id)
    T.ensure_writable(db, ctx, unit)
    errors: list[dict] = []
    changes: list[dict] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - 非 xlsx 文件统一给出可读提示
        raise ValidationFailed("文件无法读取，请上传由「下载填报模板」生成的 .xlsx 文件", code="bad_file") from None
    for sheet in _sheets(db, unit):
        if sheet.title not in wb.sheetnames:
            errors.append({"sheet": sheet.title, "row": 1, "field": "工作表", "message": f"缺少工作表「{sheet.title}」"})
            continue
        expected = _template_rows(db, ctx, unit, sheet)
        index = {(tuple(d), fmt_month(m)): (key, row) for d, key, m, row in expected}
        allowed_dims = {i: {d[i] for d, *_ in expected} for i in range(len(sheet.dims))}
        ws = wb[sheet.title]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        item_by_header = {_col_header(lbl, kind): item for item, (lbl, kind) in sheet.items.items()}
        cols: dict[int, str] = {}
        for ci, h in enumerate(header):
            if ci < len(sheet.dims) or h == "月份" or h == "":
                continue
            if h not in item_by_header:
                errors.append({"sheet": sheet.title, "row": 1, "field": h,
                               "message": f"「{h}」不是本人负责的输入项，请删除该列后重新上传"})
                continue
            cols[ci] = item_by_header[h]
        month_ci = header.index("月份") if "月份" in header else None
        for ri, values in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            dims = tuple("" if values[i] is None else str(values[i]).strip() for i in range(len(sheet.dims)))
            bad_dim = next((i for i, d in enumerate(dims) if d not in allowed_dims[i]), None)
            if bad_dim is not None:
                errors.append({"sheet": sheet.title, "row": ri, "field": sheet.dims[bad_dim],
                               "message": f"「{dims[bad_dim]}」不在你的数据范围内"})
                continue
            raw_month = values[month_ci] if month_ci is not None and month_ci < len(values) else None
            month_key = raw_month.strftime("%Y-%m") if isinstance(raw_month, date) else str(raw_month or "").strip()
            hit = index.get((dims, month_key))
            if hit is None:
                errors.append({"sheet": sheet.title, "row": ri, "field": "月份",
                               "message": f"「{month_key}」不是本次编制年度的月份"})
                continue
            key, _ = hit
            for ci, item in cols.items():
                raw = values[ci] if ci < len(values) else None
                label, kind = sheet.items[item]
                try:
                    parse_cell(kind, raw)
                except CellError as e:
                    errors.append({"sheet": sheet.title, "row": ri, "field": label, "message": e.template_message})
                    continue
                changes.append({"row_key": f"{key}|{item}", "field": month_key,
                                "value": raw if not isinstance(raw, float) else format(raw, ".15g")})
    batch = ImportBatch(source=ImportSource.TEMPLATE, file_name=file_name or "填报模板.xlsx", file_size=len(content),
                        status=ImportStatus.VALIDATION_FAILED if errors else ImportStatus.EFFECTIVE,
                        errors=errors or None, row_count=len(changes), unit_id=unit.id, uploaded_by=actor.user.id,
                        uploaded_at=clock.now(), validated_at=clock.now())
    db.add(batch)
    db.flush()
    if errors:
        write_audit(db, action="template.rejected", object_type="filling_unit", object_id=unit.id, actor=actor,
                    request=request, after={"batch_id": batch.id, "errors": len(errors)})
        return {"status": "failed", "batch_id": batch.id, "errors": errors, "written": 0}
    result = save_cells(db, ctx, actor, unit.id, changes, check_base=False, request=request, source="template")
    batch.confirmed_at = clock.now()
    batch.confirmed_by = actor.user.id
    return {"status": "imported", "batch_id": batch.id, "errors": [], "written": result["saved"]}


# ---------------------------------------------------------------------------
# 损益汇总（API-F06-01）与公司层其他损益（API-F06-02）
# ---------------------------------------------------------------------------
def version_pl(db: Session, version_id: int) -> dict:
    """按月：平台渠道毛利 + B 端毛利 − 部门费用（含专项）+ 公司层其他损益 = 公司净利润（AC-F06-07）。
    其他损益只出现在公司层，不分摊到渠道、品类或部门（AC-F06-06）。"""
    version = db.get(BudgetVersion, version_id)
    if version is None:
        raise NotFound("版本不存在")
    months = year_months(version.year)

    def by_month(stmt) -> dict[date, tuple]:
        return {r[0]: tuple(r[1:]) for r in db.execute(stmt).all()}

    pf = by_month(select(PlatformUnit.month, func.sum(PlatformUnit.revenue_cny), func.sum(PlatformUnit.gross_profit))
                  .where(PlatformUnit.version_id == version_id).group_by(PlatformUnit.month))
    bb = by_month(select(B2bUnit.month, func.sum(B2bUnit.revenue_cny), func.sum(B2bUnit.gross_profit))
                  .where(B2bUnit.version_id == version_id).group_by(B2bUnit.month))
    de = by_month(select(DeptExpense.month, func.sum(DeptExpense.amount))
                  .where(DeptExpense.version_id == version_id).group_by(DeptExpense.month))
    rs = by_month(select(Reserve.plan_month, func.sum(Reserve.amount))
                  .where(Reserve.version_id == version_id).group_by(Reserve.plan_month))
    op = by_month(select(CompanyOtherPl.month, func.sum(CompanyOtherPl.amount))
                  .where(CompanyOtherPl.version_id == version_id).group_by(CompanyOtherPl.month))
    out_months = []
    tot = defaultdict(lambda: ZERO)
    for m in months:
        prev, pgp = pf.get(m, (None, None))
        brev, bgp = bb.get(m, (None, None))
        dept = (de.get(m, (None,))[0] or ZERO) + (rs.get(m, (None,))[0] or ZERO)
        other = op.get(m, (None,))[0] or ZERO
        revenue = (prev or ZERO) + (brev or ZERO)
        gp = (pgp or ZERO) + (bgp or ZERO)
        net = company_net_profit(pgp or ZERO, bgp or ZERO, dept, other)
        vals = {"revenue_cny": revenue, "platform_gross_profit": pgp or ZERO, "b2b_gross_profit": bgp or ZERO,
                "gross_profit": gp, "dept_expense": dept, "other_pl": other, "net_profit": net}
        for k, v in vals.items():
            tot[k] += v
        out_months.append({"month": fmt_month(m), **{k: str(v) for k, v in vals.items()},
                           **{f"{k}_display": fmt_amount(v) for k, v in vals.items()}})
    channels = db.execute(select(Channel.name, func.sum(PlatformUnit.revenue_cny), func.sum(PlatformUnit.gross_profit))
                          .join(Channel, Channel.id == PlatformUnit.channel_id)
                          .where(PlatformUnit.version_id == version_id)
                          .group_by(Channel.name, Channel.sort).order_by(Channel.sort)).all()
    customers = db.execute(select(BCustomer.name, func.sum(B2bUnit.revenue_cny), func.sum(B2bUnit.gross_profit))
                           .join(BCustomer, BCustomer.id == B2bUnit.customer_id)
                           .where(B2bUnit.version_id == version_id)
                           .group_by(BCustomer.name, BCustomer.id).order_by(BCustomer.id)).all()

    def dim_rows(rows, kind):
        return [{"kind": kind, "name": n, "revenue_cny": str(r or ZERO), "gross_profit": str(g or ZERO),
                 "revenue_display": fmt_amount(r or ZERO), "gross_profit_display": fmt_amount(g or ZERO),
                 "gross_margin_display": fmt_percent(margin_of(g, r), 1)} for n, r, g in rows]

    total = {k: str(v) for k, v in tot.items()}
    total.update({f"{k}_display": fmt_amount(v) for k, v in tot.items()})
    total["gross_margin_display"] = fmt_percent(margin_of(tot["gross_profit"], tot["revenue_cny"]), 1)
    total["net_margin_display"] = fmt_percent(margin_of(tot["net_profit"], tot["revenue_cny"]), 1)
    return {"version": {"id": version.id, "name": version.name, "year": version.year, "type": version.type,
                        "locked": version.locked},
            "months": out_months, "channels": dim_rows(channels, "channel"),
            "customers": dim_rows(customers, "customer"), "total": total}


def company_other_pl(db: Session, version_id: int) -> list[dict]:
    rows = db.scalars(select(CompanyOtherPl).where(CompanyOtherPl.version_id == version_id)
                      .order_by(CompanyOtherPl.item, CompanyOtherPl.month))
    return [{"month": fmt_month(r.month), "item": r.item, "item_label": OtherPlItem.LABELS[r.item],
             "amount": str(r.amount), "amount_display": fmt_amount(r.amount)} for r in rows]


# ---------------------------------------------------------------------------
# 滚动预测更新（AC-F08-05；对接 G3 versions：plan_units / rolling_submissions / submit_rolling_unit）
# 滚动预测不生成 FillingUnit，这里按 G3 推导的单位构造临时单位对象，复用同一表格、保存与增量重算。
# ---------------------------------------------------------------------------
def rolling_unit(db: Session, version_id: int, unit_key: str) -> FillingUnit:
    vmod = versions_module()
    version = db.get(BudgetVersion, version_id)
    if vmod is None or version is None or version.type not in (VersionType.ROLLING_DRAFT, VersionType.SNAPSHOT):
        raise NotFound("滚动预测不存在")
    plan = next((p for p in vmod.plan_units(db) if p.key == unit_key), None)
    if plan is None or plan.filler is None or plan.reviewer is None:
        raise NotFound("填报单位不存在")
    submitted = unit_key in vmod.rolling_submissions(db, version_id)
    unit = FillingUnit(id=0, version_id=version_id, unit_type=plan.unit_type, scope_object_id=plan.scope_object_id,
                       name=plan.name, filler_id=plan.filler.id, reviewer_id=plan.reviewer.id,
                       status=UnitStatus.SUBMITTED if submitted or version.type != VersionType.ROLLING_DRAFT
                       else UnitStatus.FILLING, updated_at=clock.now())
    setattr(unit, "rolling_key", unit_key)  # 临时单位标识（非映射属性）
    return unit


def my_rolling_units(db: Session, ctx: ScopeContext) -> list[dict]:
    vmod = versions_module()
    if vmod is None:
        return []
    out = []
    for v in db.scalars(select(BudgetVersion).where(BudgetVersion.type == VersionType.ROLLING_DRAFT)
                        .order_by(BudgetVersion.id)):
        for p in vmod.plan_units(db):
            if p.filler is None or p.reviewer is None:
                continue
            unit = rolling_unit(db, v.id, p.key)
            if T.can_fill(db, ctx, unit):
                out.append(T.unit_info(db, ctx, unit))
    return out


def submit_rolling(db: Session, ctx: ScopeContext, actor: CurrentUser, version_id: int, unit_key: str) -> dict:
    unit = rolling_unit(db, version_id, unit_key)
    T.ensure_writable(db, ctx, unit)
    vmod = versions_module()
    vmod.submit_rolling_unit(db, vmod.as_actor(db, actor), version_id, unit_key)
    return T.unit_info(db, ctx, rolling_unit(db, version_id, unit_key))
