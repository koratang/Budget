"""REQ-F09-03 经营历史表导入、REQ-F04-06 填报上年同月参考、REQ-F10-07 偏差报告同比。

样例表按财务部经营分析表结构构造（「<渠道>-RMB」按店铺分块、「管理费用RMB」取「N月实际」）；
期望值按毛利公式手算：毛利 = 收入 − 退款 − 采购成本 − 头程 − 尾程 − 关税 − 佣金 − 推广 − 仓储 − 其他运营费用 + 库存索赔。
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func, select

from app.models import BudgetAccount, KingdeeActual, Store
from tests.acceptance.e3_scenarios import flow32

D = Decimal
ITEMS = ["收入", "售后退款", "商品成本", "头程物流", "尾程物流", "关税成本", "平台佣金", "推广费用", "仓储费用",
         "运营费用（软件费用等）", "库存索赔金额", "毛利额（RMB）"]
# 每月：收入, 退款, 采购, 头程, 尾程, 关税, 佣金, 推广, 仓储, 其他运营, 库存索赔（表内负数）
JAN_2026 = [100000, 10000, 30000, 5000, 8000, 0, 15000, 12000, 2000, 1000, -500]   # 毛利 17,500
FEB_2026 = [80000, 8000, 24000, 4000, 6400, 0, 12000, 9600, 1600, 800, -400]       # 毛利 14,000
JAN_2027 = [120000, 12000, 36000, 6000, 9600, 0, 18000, 18400, 2400, 1200, -600]   # 毛利 17,000


def gp(v: list[int]) -> int:
    return v[0] - sum(v[1:10]) - v[10]


def workbook(blocks: dict[str, list[tuple[str, dict[int, list[int]], dict[int, int] | None]]],
             mgmt: list[tuple[str, str, str, dict[int, int]]]) -> bytes:
    """blocks：工作表 → [(分块名, {月: 各项}, {月: 表内毛利额} 或 None 按公式)]；mgmt：[(项目, 说明, 负责部门, {月: 金额})]。"""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, items in blocks.items():
        ws = wb.create_sheet(sheet)
        ws.append([None, "项目", "说明", "年度合计"] + [f"{m}月" for m in range(1, 13)])
        ws.append(["汇总", "收入", None, 999999999] + [999999999] * 12)      # 汇总分块不导入
        for name, months, sheet_gp in items:
            for i, label in enumerate(ITEMS):
                row = [name if i == 0 else None, label, None, None]
                for m in range(1, 13):
                    vals = months.get(m)
                    if vals is None:
                        row.append(0)
                    elif label.startswith("毛利额"):
                        row.append((sheet_gp or {}).get(m, gp(vals)))
                    else:
                        row.append(vals[i])
                ws.append(row)
        ws.append(["毛利额验证", "毛利额", None, 1] + [1] * 12)             # 验证分块不导入
    ws = wb.create_sheet("管理费用RMB")
    ws.append(["项目", "说明", "负责部门"] + [f"{m}月实际" for m in range(1, 10)] + ["10月预算", "11月预算", "12月预算"])
    for item, note, dept, months in mgmt:
        ws.append([item, note, dept] + [months.get(m) for m in range(1, 10)] + [88888, 88888, 88888])
    ws.append(["管理费用合计", None, None] + [1] * 12)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


WB_2026 = dict(
    blocks={"Amazon-RMB": [("Amazon-US-01", {1: JAN_2026, 2: FEB_2026}, {2: 13000}), ("Amazon-CA-02", {1: JAN_2026}, None)],
            "TO B-RMB": [("Ferguson", {1: [50000, 0, 20000, 0, 0, 0, 0, 0, 0, 0, 0]}, None)]},
    mgmt=[("招聘费", "", "人力资源部", {1: 5000}), ("福利费", "节日福利", "行政部", {1: 3000}),
          ("财务费用", "汇兑损益", "财务部", {1: -200})])
WB_2027 = dict(blocks={"Amazon-RMB": [("Amazon-US-01", {1: JAN_2027}, None)]},
               mgmt=[("招聘费", "", "人力资源部", {1: 6000})])


def upload(g3, content: bytes, year: int | None, name: str = "2026年经营分析.xlsx") -> dict:
    c = g3.client("陈思雨")
    data = {"source": "ops-history"}
    if year is not None:
        data["year"] = str(year)
    r = c.post("/api/imports", files={"file": (name, content, "application/vnd.ms-excel")}, data=data)
    if r.status_code != 200:
        return r.json()
    r2 = c.get(f"/api/imports/{r.json()['id']}")
    assert r2.status_code == 200, r2.text
    return r2.json()


def import_year(g3, wb: dict, year: int, overwrite: bool = False) -> dict:
    batch = upload(g3, workbook(**wb), year, f"{year}年经营分析.xlsx")
    if batch["status"] == "failed":
        r = g3.client("陈思雨").post(f"/api/imports/{batch['id']}/create-missing")
        assert r.status_code == 200, r.text
        batch = r.json()
    assert batch["status"] == "passed", batch
    r = g3.confirm(batch["id"], overwrite=overwrite)
    assert r.status_code == 200, r.text
    return r.json()


def test_year_is_required(g3):
    r = upload(g3, workbook(**WB_2026), None)
    assert r["code"] == "year_required"
    assert r["message"] == "请选择表格所属年度"


def test_missing_master_plan_create_and_confirm(g3):
    db = g3.db
    batch = upload(g3, workbook(**WB_2026), 2026)
    assert batch["status"] == "failed"
    blocking = {(i["code"], i["name"]) for i in batch["issues"] if not i["warning"]}
    assert blocking == {("store_not_found", "Amazon-CA-02（Amazon-RMB）"), ("mgmt_account_not_found", "福利费")}
    plan = batch["missing_master"]
    assert plan["stores"] == [{"name": "Amazon-CA-02", "channel": "Amazon", "new_channel": False, "site": "加拿大",
                               "site_code": "CA", "currency": "CAD"}]
    assert plan["accounts"] == [{"name": "福利费"}]
    assert plan["count"] == 2

    r = g3.client("陈思雨").post(f"/api/imports/{batch['id']}/create-missing")
    assert r.status_code == 200, r.text
    batch = r.json()
    assert batch["created"] == {"stores": 1, "channels": 0, "customers": 0, "accounts": 1}
    assert batch["status"] == "passed"
    assert batch["periods"] == ["2026-01", "2026-02"]
    warnings = {(i["code"], i["name"]) for i in batch["issues"] if i["warning"]}
    assert ("ops_gp_mismatch", "Amazon-RMB · Amazon-US-01 · 2月：重算 ¥ 14,000.00，表内 ¥ 13,000.00") in warnings
    assert ("ops_dept_public", "行政部") in warnings
    new_store = db.scalar(select(Store).where(Store.name == "Amazon-CA-02"))
    assert new_store.currency == "CAD"
    assert db.scalar(select(BudgetAccount).where(BudgetAccount.name == "福利费")).is_dept_expense is True

    r = g3.confirm(batch["id"])
    assert r.status_code == 200, r.text
    us = db.scalar(select(Store).where(Store.name == "Amazon-US-01"))

    def store_sum(item: str, month: str) -> Decimal:
        return db.scalar(select(func.sum(KingdeeActual.amount)).join(BudgetAccount)
                         .where(KingdeeActual.store_id == us.id, BudgetAccount.pl_item == item,
                                KingdeeActual.month == date.fromisoformat(f"{month}-01")))

    assert store_sum("revenue", "2026-01") == D("100000")
    assert store_sum("inventory_claim", "2026-01") == D("500")       # 表内 −500 冲减成本 → 正向项目
    assert store_sum("refund", "2026-02") == D("8000")
    # 3 月全部为 0：不入库；汇总与验证分块不入库
    assert db.scalar(select(func.count()).select_from(KingdeeActual)
                     .where(KingdeeActual.month == date(2026, 3, 1))) == 0
    fx = db.scalar(select(KingdeeActual).join(BudgetAccount).where(BudgetAccount.name == "汇兑损益"))
    assert fx.amount == D("-200")                                    # 「财务费用 / 汇兑损益」按说明归到汇兑损益
    welfare = db.scalar(select(KingdeeActual).join(BudgetAccount).where(BudgetAccount.name == "福利费"))
    assert welfare.department_id is None and welfare.amount == D("3000")   # 行政部不在主数据 → 公司公共


def test_reimport_same_period_needs_overwrite(g3):
    import_year(g3, WB_2026, 2026)
    batch = upload(g3, workbook(**WB_2026), 2026)
    assert batch["status"] == "passed" and batch["needs_overwrite"] is True
    r = g3.confirm(batch["id"])
    assert r.status_code == 409 and r.json()["code"] == "needs_overwrite"


def test_yoy_in_channel_site_report(g3):
    import_year(g3, WB_2026, 2026)
    import_year(g3, WB_2027, 2027)
    c = g3.client("林晓")
    rep = c.get("/api/reports/channel-site", params={"month": "2027-01"}).json()
    assert rep["prior_month"] == "2026-01" and rep["prior_available"] is True
    us = next(r for r in rep["rows"] if r["name"] == "Amazon-US-01")
    assert D(us["yoy"]["revenue"]["prior"]) == D("100000")
    assert D(us["yoy"]["revenue"]["growth"]) == D("20")               # (120,000 − 100,000) ÷ 100,000
    assert D(us["yoy"]["gross_profit"]["prior"]) == D("17500")
    assert round(D(us["yoy"]["gross_profit"]["growth"]), 2) == D("-2.86")   # (17,000 − 17,500) ÷ 17,500
    ca = next(r for r in rep["rows"] if r["name"] == "Amazon-CA-02")      # 今年无实际、上年有：仍列出
    assert D(ca["yoy"]["revenue"]["growth"]) == D("-100")


def test_yoy_unavailable_without_prior_year(g3):
    import_year(g3, WB_2027, 2027)
    rep = g3.client("林晓").get("/api/reports/channel-site", params={"month": "2027-01"}).json()
    assert rep["prior_available"] is False
    us = next(r for r in rep["rows"] if r["name"] == "Amazon-US-01")
    assert us["yoy"]["available"] is False and us["yoy"]["revenue"]["growth"] is None


def test_platform_grid_shows_prior_year_reference_and_export(g3):
    _versions, unit = flow32(g3.db)
    import_year(g3, WB_2026, 2026)
    c = g3.client("林晓")
    grid = c.get(f"/api/filling-units/{unit.id}/grid").json()
    refs = [r for r in grid["rows"] if r["kind"] == "ref"]
    assert grid["rows"][:len(refs)] == refs                           # 参考行在表格最上方
    by = {r["values"]["item"]: r["values"] for r in refs}
    assert by["收入（人民币）"]["2027-01"] == "100,000.00"              # 2026-01 实际放在 2027-01 列
    assert by["收入（人民币）"]["2027-02"] == "80,000.00"
    assert by["收入（人民币）"]["2027-03"] == "—"
    assert by["收入（人民币）"]["year_total"] == "180,000.00"
    assert by["毛利（人民币）"]["2027-01"] == "17,500.00"
    assert by["毛利率"]["2027-01"] == "17.50%"
    assert by["退款率"]["2027-02"] == "10.00%"                         # 8,000 ÷ 80,000
    assert by["推广费用率"]["2027-01"] == "12.00%"
    assert refs[0]["values"]["category"] == "2026 实际"
    assert grid["prior"] == {"year": 2026, "months": 2, "span": "1–2 月", "revenue_display": "¥ 180,000.00",
                             "gross_profit_display": "¥ 31,500.00", "gross_margin_display": "17.5%"}
    # 导出与表格一致：参考行同样在最上方
    r = c.get("/api/export", params={"view": "budget-platform", "filters": f'{{"unit_id": {unit.id}}}'})
    assert r.status_code == 200, r.text
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(r.content)).active
    first = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    assert first[:4] == ["2026 实际", "参考", "收入（人民币）", D("100000.00")]
