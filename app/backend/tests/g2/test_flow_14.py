"""FLOW-14：Excel 模板下载与上传（AC-F04-07/08）。

契约期望：王浩模板只含 Amazon-US-01 的维度组合与 7 个驱动列，已保存值（如 FLOW-11 的 SPU 10、单价 80.00）已预填；
蒋帆模板只含四项比例列，无收入、销量、利润列；上传后页面数据与数据库值均未变化，错误清单列出 3 条，
各含行号、字段与原因；修正后上传成功并刷新合计
"""
import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import PlatformUnit
from tests.g2.helpers import grid, login, minimal_units, publish, row, unit_by_name
from tests.g2.test_flow_11 import fill_flow11_drivers

DRIVER_HEADERS = ["SPU 数量", "单 SPU 平均销量", "平均单价（USD）", "折扣率（%）", "退款率（%）", "广告费率（%）",
                  "佣金率（%）"]


def _download(client, unit_id):
    r = client.get(f"/api/filling-units/{unit_id}/template")
    assert r.status_code == 200, r.text
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content))


def _upload(client, unit_id, wb):
    buf = io.BytesIO()
    wb.save(buf)
    return client.post(f"/api/filling-units/{unit_id}/template-upload",
                       files={"file": ("Amazon-US-01_2027年度预算填报模板.xlsx", buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _snapshot(db, store_id):
    db.expire_all()
    return sorted((u.id, u.spu_count, u.avg_sales, u.avg_price, u.discount_rate, u.revenue_cny)
                  for u in db.scalars(select(PlatformUnit).where(PlatformUnit.store_id == store_id)))


def test_flow_14_template_download_and_upload(seeded, client):
    db = seeded
    task_id = publish(client, db, minimal_units(db)).json()["task"]["id"]
    unit = unit_by_name(db, task_id, "Amazon-US-01")
    fill_flow11_drivers(client, db, unit.id)

    login(client, db, "wanghao")
    wb = _download(client, unit.id)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert header == ["店铺", "站点", "品类", "新老品", "月份", *DRIVER_HEADERS]
    assert {r[0] for r in rows[1:]} == {"Amazon-US-01"}
    hit = next(r for r in rows[1:] if r[2] == "浴室柜" and r[3] == "新品" and r[4] == "2027-03")
    assert hit[5] == "10" and hit[7] == "80.00"

    login(client, db, "jiangfan")
    lwb = _download(client, unit_by_name(db, task_id, "物流比例（物流物控）").id)
    for name in lwb.sheetnames:
        h = [c for c in next(lwb[name].iter_rows(values_only=True))]
        assert [x for x in h if "率" in str(x)] == ["头程费率（%）", "尾程费率（%）", "仓储费率（%）", "关税率（%）"]
        assert not any(w in str(x) for x in h for w in ("收入", "销量", "利润"))

    # 第 5 行 SPU 数量 −1、第 9 行折扣率 130、第 12 行店铺「Amazon-JP-01」，其余行合法
    login(client, db, "wanghao")
    before = _snapshot(db, unit.scope_object_id)
    ws["F5"] = "-1"
    ws["I9"] = "130"
    ws["A12"] = "Amazon-JP-01"
    ws["F3"] = "7"
    ws["G3"] = "5"
    r = _upload(client, unit.id, wb)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "failed" and body["written"] == 0
    errs = sorted((e["row"], e["field"]) for e in body["errors"])
    assert errs == [(5, "SPU 数量"), (9, "折扣率"), (12, "店铺")]
    assert all(e["message"] for e in body["errors"])
    assert "Amazon-JP-01" in next(e["message"] for e in body["errors"] if e["row"] == 12)
    assert _snapshot(db, unit.scope_object_id) == before

    # 修正后上传成功并刷新合计
    g0 = grid(client, unit.id)
    ws["F5"] = "0"
    ws["I9"] = "0"
    ws["A12"] = "Amazon-US-01"
    r = _upload(client, unit.id, wb)
    assert r.json()["status"] == "imported", r.text
    g1 = grid(client, unit.id)
    assert row(g1, "total|sales_qty")["values"] != row(g0, "total|sales_qty")["values"]
    assert row(g1, "total|revenue_cny")["values"]["2027-03"] == row(g0, "total|revenue_cny")["values"]["2027-03"]
