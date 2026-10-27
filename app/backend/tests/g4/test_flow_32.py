"""FLOW-32 Excel 导出：表头与页面当前显示列一致、只含用户范围内的行、金额与页面显示值逐格相同（2 位小数）。

期望值逐字取自 TDD验收契约.md FLOW-32 行。
"""
import io
import json
from datetime import date
from decimal import Decimal
from urllib.parse import unquote

from openpyxl import load_workbook

from app.core.money import round2
from app.models import BudgetTask, FillingUnit
from app.models.enums import UnitStatus, UnitType

from tests.g4 import builders as B
from tests.g4.scenarios import flow25

# PAGE-F10-02 §5.2 主表列（达成条列以数据条导出）
CHANNEL_GROUPS = ["店铺", "渠道", "站点", "收入", "收入差额拆分", "毛利", "毛利率", "偏差说明"]
CHANNEL_HEADERS = ["店铺", "渠道", "站点", "预算", "实际", "实际（本币）", "差额", "达成率", "达成条",
                   "量差", "价差", "汇率差", "口径差", "预算", "实际", "差额", "达成率", "预算", "实际", "差（百分点）",
                   "偏差说明"]
SPLIT = {"量差", "价差", "汇率差", "口径差"}


def export(client, view, filters):
    r = client.get("/api/export", params={"view": view, "filters": json.dumps(filters)})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    name = unquote(r.headers["content-disposition"].split("filename*=UTF-8''")[1])
    assert name.endswith(".xlsx")
    return name, load_workbook(io.BytesIO(r.content))


def header_row(ws, row):
    return [c.value for c in ws[row]]


def money_cell(cell) -> Decimal:
    assert cell.number_format == "#,##0.00"
    v = Decimal(str(cell.value))
    assert v == round2(v)  # 保留 2 位小数
    return v


def first_col(ws, start):
    return [r[0].value for r in ws.iter_rows(min_row=start)]


def test_lin_xiao_exports_all_reports(client, seeded):
    db = seeded
    flow25(db)
    B.login(client, db, "linxiao")

    # 渠道×站点偏差：两行表头与页面一致；含全部店铺、B 端与未归属渠道
    page = client.get("/api/reports/channel-site", params={"month": "2027-04"}).json()
    name, wb = export(client, "report-channel", {"month": "2027-04", "baseline": "adj", "channel": "all"})
    assert name == "渠道站点偏差_2027-04_最新调整后目标.xlsx"
    ws = wb["渠道站点偏差"]
    assert [v for v in header_row(ws, 1) if v] == CHANNEL_GROUPS
    assert header_row(ws, 2) == CHANNEL_HEADERS
    names = first_col(ws, 3)
    assert names == ["Amazon-US-01", "Amazon-CA-01", "B 端（按客户）", "Ferguson", "未归属渠道",
                     "合计（含 B 端与未归属渠道）"]
    # 导出金额与页面显示值逐格相同、保留 2 位小数
    for row in ws.iter_rows(min_row=3):
        label = row[0].value
        src = next((r for r in page["rows"] + page["b2b_rows"] if r["name"] == label), None)
        if src is None:
            continue
        expect = [src["revenue"]["budget"], src["revenue"]["actual"], None, src["revenue"]["diff"]]
        for cell, v in zip(row[3:7], expect):
            if v is not None:
                assert money_cell(cell) == round2(Decimal(v))
        assert money_cell(row[13]) == round2(Decimal(src["gross_profit"]["budget"]))
        assert money_cell(row[14]) == round2(Decimal(src["gross_profit"]["actual"]))
    us = next(r for r in ws.iter_rows(min_row=3) if r[0].value == "Amazon-US-01")
    assert money_cell(us[3]) == Decimal("615168.00")
    assert money_cell(us[6]) == Decimal("-60768.00")
    assert us[7].value == Decimal("0.9012") or abs(Decimal(str(us[7].value)) - Decimal("0.9012")) < Decimal("1e-9")
    assert us[7].number_format == "0.00%"
    assert [money_cell(c) for c in us[9:13]] == [Decimal("-51264.00"), Decimal("-15664.00"), Decimal("6160.00"),
                                                Decimal("0.00")]
    detail = wb["明细"]
    promo = next(r for r in detail.iter_rows(min_row=2) if r[0].value == "Amazon-US-01" and r[1].value == "推广费用")
    assert money_cell(promo[9]) == Decimal("6472.00")

    # BI 未到的月份：隐藏拆分列与费率三列时同样不导出
    _, wb5 = export(client, "report-channel", {"month": "2027-05"})
    assert header_row(wb5["渠道站点偏差"], 2) == [h for h in CHANNEL_HEADERS if h not in SPLIT]
    assert "费率差" not in header_row(wb5["明细"], 1)

    # 公司汇总：三个工作表
    name, wb = export(client, "report-company", {})
    assert name == "公司汇总_截至2027-05.xlsx"
    assert wb.sheetnames == ["净利润", "增速对比", "预测准确度"]
    assert header_row(wb["净利润"], 1) == ["月份", "年初原始目标净利润", "最新调整后目标净利润", "实际净利润",
                                         "达成率（原始）", "达成率（调整后）", "达成条（调整后）"]

    # 品类报告
    name, wb = export(client, "report-category", {"month": "2027-04"})
    assert name == "品类新老品偏差_2027-04.xlsx"
    ws = wb.active
    assert header_row(ws, 2) == ["品类", "新品/老品"] + ["预算", "实际", "差额"] * 4 + ["偏差说明"]
    bath = next(r for r in ws.iter_rows(min_row=3) if r[0].value == "浴室柜" and r[1].value == "新品")
    assert [money_cell(c) for c in bath[8:11]] == [Decimal("61516.80"), Decimal("61600.00"), Decimal("83.20")]


def test_wang_hao_export_only_own_store(client, seeded):
    db = seeded
    flow25(db)
    B.login(client, db, "wanghao")
    _, wb = export(client, "report-channel", {"month": "2027-04"})
    ws = wb["渠道站点偏差"]
    assert first_col(ws, 3) == ["Amazon-US-01", "合计（当前范围）"]
    detail_names = {r[0].value for r in wb["明细"].iter_rows(min_row=2)}
    assert detail_names == {"Amazon-US-01"}
    # 公司汇总对运营无权限
    r = client.get("/api/export", params={"view": "report-company", "filters": "{}"})
    assert r.status_code == 403 and r.json() == {"code": "forbidden", "message": "无权限"}


def test_dept_and_budget_table_export(client, seeded):
    db = seeded
    versions = flow25(db)
    target = versions["target"]
    for m in range(1, 13):
        B.dept_expense(db, target, "人力资源部", "培训费", f"2027-{m:02d}", "100000.00")
    kd = B.batch(db, "kingdee", "2027-01")
    B.kingdee(db, kd, "2027-01", "培训费", "123456.785", dept_name="人力资源部")
    task = BudgetTask(name="2027 年度预算编制", year=2027, deadline=date(2026, 12, 20), version_id=target.id)
    db.add(task)
    db.flush()
    us = B.store(db, "Amazon-US-01")
    unit = FillingUnit(task_id=task.id, version_id=target.id, unit_type=UnitType.PLATFORM, scope_object_id=us.id,
                       name="Amazon-US-01", filler_id=B.user_id(db, "wanghao"), reviewer_id=B.user_id(db, "zhaolei"),
                       status=UnitStatus.SUBMITTED)
    db.add(unit)
    db.flush()

    B.login(client, db, "linxiao")
    name, wb = export(client, "report-dept", {"year": 2027})
    assert name == "部门费用报告_截至2027-05.xlsx"
    ws = wb.active
    assert header_row(ws, 1) == ["部门", "预算科目", "全年预算", "累计实际", "额度使用率", "时间进度", "达成条", "偏差说明"]
    row = next(r for r in ws.iter_rows(min_row=2) if r[1].value == "培训费")
    assert money_cell(row[2]) == Decimal("1200000.00")
    assert money_cell(row[3]) == Decimal("123456.79")  # ROUND_HALF_UP 2 位，与页面显示一致

    name, wb = export(client, "budget-platform", {"unit_id": unit.id})
    assert name == "预算表_Amazon-US-01_2027 年度预算编制.xlsx"
    ws = wb.active
    assert header_row(ws, 1) == ["品类", "新/老品", "项目"] + [f"{m} 月" for m in range(1, 13)] + ["全年"]
    rev = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "浴室柜" and r[2].value == "收入（人民币）")
    assert money_cell(rev[3 + 3]) == Decimal("615168.00")  # 4 月

    # 其他店铺的运营（李佳琪）导出 Amazon-US-01 预算表：403
    B.login(client, db, "lijiaqi")
    r = client.get("/api/export", params={"view": "budget-platform", "filters": json.dumps({"unit_id": unit.id})})
    assert r.status_code == 403
