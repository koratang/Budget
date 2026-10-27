"""E2 浏览器验收前置（库 budget_acc2）：重建库 + 种子 + 契约行写明的初始状态，结果以 JSON 打印给 Playwright 脚本。

前置经真实路由（进程内 TestClient，提交到 budget_acc2）完成，只做契约「初始状态」，待验收动作留给浏览器步骤。
用法（app/backend 下）：
    .venv/bin/python -m tests.acceptance.e2_setup reset <flow_11 … flow_20>
    .venv/bin/python -m tests.acceptance.e2_setup step <名称> [参数…]      # 场景中途由契约写明的外部前置（如金蝶 4 月实际导入）
    .venv/bin/python -m tests.acceptance.e2_setup dump <version_id>        # 版本全部单元值（用于调整前后逐值比对）
"""
import os
import sys

URL = "postgresql+psycopg://budget:budget@127.0.0.1:5432/budget_acc2"
os.environ.update(DATABASE_URL=URL, APP_ENV="dev", DEV_IDENTITY_SWITCH="true", SCHEDULER_ENABLED="false",
                  DINGTALK_BASE_URL="http://127.0.0.1:8022", DINGTALK_APP_KEY="stub", DINGTALK_APP_SECRET="stub",
                  DINGTALK_AGENT_ID="10001", DINGTALK_CORP_ID="ding-stub", APP_BASE_URL="http://127.0.0.1:8012")

import hashlib  # noqa: E402
import json  # noqa: E402

from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import create_engine, select  # noqa: E402

from app.config import BACKEND_DIR, load_settings  # noqa: E402
from app.db import get_db, session_factory  # noqa: E402
from app.main import create_app  # noqa: E402
from app.models import BCustomer, BudgetAccount, BudgetVersion, Category, Department, DeptExpense, PlatformUnit, User  # noqa: E402
from app.models.enums import Role  # noqa: E402
from scripts.seed import seed  # noqa: E402
from tests.acceptance.e2_common import Ctx, actual_rows_2027  # noqa: E402
from tests.g2.helpers import add_user, minimal_units, pu_key, store  # noqa: E402


def rebuild() -> None:
    eng = create_engine(URL)
    with eng.begin() as conn:
        conn.exec_driver_sql("DROP SCHEMA IF EXISTS public CASCADE")
        conn.exec_driver_sql("CREATE SCHEMA public")
    eng.dispose()
    cfg = Config(str(BACKEND_DIR / "alembic.ini"))
    cfg.attributes["url"] = URL
    command.upgrade(cfg, "head")


def context():
    db = session_factory(URL)()
    app = create_app(load_settings(DATA_DIR=str(BACKEND_DIR / "data" / "e2")))

    def _db():
        yield db

    app.dependency_overrides[get_db] = _db
    return Ctx(db, TestClient(app, base_url="http://testserver"))


def scenario(name: str, ctx: Ctx) -> dict:
    db = ctx.db
    out: dict = {}
    if name == "flow_11":
        task = ctx.publish()
        ctx.preset_flow11_ratios(task)
        out["unit"] = ctx.unit(task, "Amazon-US-01")
    elif name == "flow_12":
        task = ctx.publish()
        ctx.preset_ferguson_ratios(task)
        out["unit"] = ctx.unit(task, "B端 · 美国业务部")
    elif name == "flow_13":
        task = ctx.publish()
        ctx.fill_flow11_drivers(task)
        ctx.preset_ferguson_ratios(task)
        ctx.fill_flow12_sales(task)
        ctx.fill_hr_300k(task)
        out.update(task=task, cost=ctx.unit(task, "成本比例（财务BP）"), logistics=ctx.unit(task, "物流比例（物流物控）"),
                   platform=ctx.unit(task, "Amazon-US-01"))
    elif name == "flow_14":
        task = ctx.publish()
        ctx.fill_flow11_drivers(task)
        out.update(unit=ctx.unit(task, "Amazon-US-01"), logistics=ctx.unit(task, "物流比例（物流物控）"),
                   store_id=store(db, "Amazon-US-01").id)
    elif name == "flow_15":
        add_user(db, "liuchang", "刘畅", "运营部", [Role.OPERATIONS], {"store": [store(db, "Amazon-US-01").id]})
        db.commit()
        task = ctx.publish()
        out["unit"] = ctx.unit(task, "Amazon-US-01")
    elif name == "flow_16":
        from tests.acceptance.test_flow_16 import prepare
        task = prepare(ctx)
        out["unit"] = ctx.unit(task, "人力资源部")
    elif name == "flow_17":
        from tests.g2.test_flow_17 import setup_rule_and_store
        setup_rule_and_store(db)
        db.commit()
        task = ctx.publish(minimal_units(db, stores=("Amazon-US-01", "TikTok-US-01")))
        out["unit"] = ctx.unit(task, "TikTok-US-01")
    elif name == "flow_18":
        task = ctx.publish()
        out.update(task=task, unit=ctx.unit(task, "Amazon-US-01"))
    elif name == "flow_19":
        from tests.acceptance.test_flow_19 import prepare
        out.update(prepare(ctx))
    elif name == "flow_20":
        out["v1"] = ctx.build_v1()
    else:
        raise SystemExit(f"未知场景 {name}")
    out["store_ids"] = {s: store(db, s).id for s in ("Amazon-US-01", "TikTok-US-01")}
    out["users"] = {u.dingtalk_userid: u.id for u in db.scalars(select(User))}
    out["names"] = {u.name: u.dingtalk_userid for u in db.scalars(select(User))}
    out["customers"] = {c.name: c.id for c in db.scalars(select(BCustomer))}
    out["accounts"] = {a.name: a.id for a in db.scalars(select(BudgetAccount))}
    out["departments"] = {d.name: d.id for d in db.scalars(select(Department))}
    amz = store(db, "Amazon-US-01")
    out["amz"] = {"channel": amz.channel_id, "site": amz.site_id,
                  "cats": {c.name: c.id for c in db.scalars(select(Category))}}
    out["keys"] = {f"{cat}|{'new' if new else 'old'}|{item}": pu_key(db, "Amazon-US-01", cat, new, item)
                   for cat in ("浴室柜", "灯镜") for new in (True, False)
                   for item in ("spu_count", "avg_sales", "avg_price", "discount_rate")}
    return out


def step(name: str, args: list[str], ctx: Ctx) -> dict:
    if name == "import-q1":  # FLOW-20：2027-04-10 已导入 1–3 月金蝶实际
        ctx.now("2027-04-10 10:00")
        return ctx.import_kingdee(actual_rows_2027([1, 2, 3]), "金蝶_2027-01至03_利润表.xlsx")
    if name == "import-april":  # FLOW-20：2027-05-10 金蝶 2027-04 实际已导入（Amazon-US-01 实际收入 ¥ 554,400.00）
        ctx.now("2027-05-10 16:00")
        return ctx.import_kingdee(actual_rows_2027([4], amz_april="554400"), "金蝶_2027-04_利润表.xlsx")
    raise SystemExit(f"未知步骤 {name}")


def dump(version_id: int) -> dict:
    with session_factory(URL)() as db:
        pu = [(u.store_id, u.category_id, u.is_new, str(u.month), u.spu_count, str(u.avg_sales), str(u.avg_price),
               str(u.revenue_cny), str(u.gross_profit))
              for u in db.scalars(select(PlatformUnit).where(PlatformUnit.version_id == version_id)
                                  .order_by(PlatformUnit.id))]
        de = [(d.department_id, d.account_id, str(d.month), str(d.amount))
              for d in db.scalars(select(DeptExpense).where(DeptExpense.version_id == version_id)
                                  .order_by(DeptExpense.id))]
        v = db.get(BudgetVersion, version_id)
        text = json.dumps([pu, de])
        return {"id": version_id, "locked": v.locked, "rows": len(pu) + len(de),
                "sha256": hashlib.sha256(text.encode()).hexdigest()}


def dump_store(store_id: int) -> dict:
    with session_factory(URL)() as db:
        rows = [(u.id, u.category_id, u.is_new, str(u.month), u.spu_count, str(u.avg_sales), str(u.avg_price),
                 str(u.discount_rate), str(u.revenue_cny))
                for u in db.scalars(select(PlatformUnit).where(PlatformUnit.store_id == store_id)
                                    .order_by(PlatformUnit.id))]
        return {"rows": len(rows), "sha256": hashlib.sha256(json.dumps(rows).encode()).hexdigest()}


def xlsx(path: str, edits: list[str]) -> dict:
    """读取（并可按 A1=值 修改后另存）下载的模板，供浏览器脚本解析 .xlsx。"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if edits:
        ws = wb[wb.sheetnames[0]]
        for e in edits:
            ref, val = e.split("=", 1)
            ws[ref] = val
        out = path.replace(".xlsx", f"_edit{len(edits)}_{abs(hash(tuple(edits))) % 10**6}.xlsx")
        wb.save(out)
        return {"saved": out}
    return {"sheets": {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}}


def main(argv: list[str]) -> None:
    cmd = argv[0]
    if cmd == "reset":
        rebuild()
        ctx = context()
        seed(ctx.db)
        result = scenario(argv[1], ctx)
        ctx.db.commit()
    elif cmd == "step":
        ctx = context()
        result = step(argv[1], argv[2:], ctx)
        ctx.db.commit()
    elif cmd == "dump":
        result = dump(int(argv[1]))
    elif cmd == "dump-store":
        result = dump_store(int(argv[1]))
    elif cmd == "xlsx":
        result = xlsx(argv[1], argv[2:])
    else:
        raise SystemExit(__doc__)
    print("E2JSON " + json.dumps(result, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main(sys.argv[1:])
