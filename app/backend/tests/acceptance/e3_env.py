"""E3 浏览器验收环境控制（库 budget_acc3 / 后端 8013 / 钉钉替身 8023）。

只写契约行的初始状态（主数据、人员与该行写明的前置数据），待验收动作由浏览器完成。
用法（app/backend 下）：
    .venv/bin/python -m tests.acceptance.e3_env reset <场景>   # 重建库 + 种子 + 场景前置（提交）
    .venv/bin/python -m tests.acceptance.e3_env add <场景>     # 在现有库上追加场景数据（不重建）
    .venv/bin/python -m tests.acceptance.e3_env serve           # 以 E3_NOW 固定时钟启动后端（托管 E3_STATIC_DIR）
    .venv/bin/python -m tests.acceptance.e3_env xlsx <文件>     # 把 .xlsx 解析为 JSON（值 + 数字格式）打印
场景生成的上传文件写到 E3_DATA_DIR（默认 /tmp/e3_acc），并打印 JSON 结果。
"""
import json
import os
import sys
from datetime import datetime
from pathlib import Path

URL = os.environ.get("E3_DATABASE_URL", "postgresql+psycopg://budget:budget@127.0.0.1:5432/budget_acc3")
DATA = Path(os.environ.get("E3_DATA_DIR", "/tmp/e3_acc"))


def _at(text: str) -> datetime:
    from app.core.clock import TZ
    return datetime.strptime(text, "%Y-%m-%d %H:%M").replace(tzinfo=TZ)


def _g3(db):
    from tests.g3.conftest import G3
    DATA.mkdir(parents=True, exist_ok=True)
    return G3(db, None, DATA)


def scenario(db, name: str) -> dict:
    from app.core.clock import clock
    from tests.acceptance import e3_scenarios as S
    from tests.g3.fixtures_g3 import fixture_path
    g3 = _g3(db)
    out: dict = {}
    if name == "flow21":
        v1, v2 = S.flow21(g3)
        out = {"v1": v1.id, "v2": v2.id}
    elif name == "flow22":
        g3.set_now("2027-01-05 09:00")
        S.flow21(g3)
        g3.set_now("2027-05-10 16:42")
        g3.import_file(fixture_path("金蝶_2027-04_利润表.xlsx"), "kingdee")
        out = {"v2": str(fixture_path("金蝶_2027-04_利润表_v2.xlsx")),
               "v3": str(fixture_path("金蝶_2027-04_利润表_v3.xlsx")),
               "history": str(fixture_path("金蝶_历史_2025-01至2026-12.xlsx"))}
    elif name == "flow23":
        g3.add_skus()
        g3.set_now("2027-04-20 09:30")
        g3.import_file(fixture_path("BI_2027-03_销售明细.xlsx"), "bi")
        out = {"bad": str(fixture_path("BI_2027-03至04_销售明细.xlsx")),
               "fixed": str(fixture_path("BI_2027-02至03_销售明细_v2.xlsx"))}
    elif name == "flow24":
        from decimal import Decimal
        from datetime import date
        from app.models import FxActual
        g3.add_skus()
        g3.set_now("2027-05-20 10:00")
        db.add(FxActual(month=date(2027, 4, 1), currency="USD", rate=Decimal("7.2000")))
        db.flush()
        g3.import_file(fixture_path("BI_2027-04_销售明细.xlsx"), "bi")
        g3.import_file(fixture_path("金蝶_2027-04_利润表_v3.xlsx"), "kingdee")
    elif name == "flow25":
        S.flow25(db)
    elif name == "flow26":
        # 浏览器场景合并品类与部门费用前置：金蝶 2027-01～04 已导入（时间进度 4 ÷ 12），BI 只有 2027-04
        v = S.flow26_category(db, with_may_kingdee=False)
        S.flow26_dept(db, v)
    elif name == "flow27":
        S.flow27(db)
    elif name == "flow27-2028":
        S.flow27_2028(db)
    elif name == "compiling-2029":
        # 工作台页头「预算年度」回归：存在进行中的 2029 年度编制任务
        from datetime import date
        from app.models import BudgetTask
        from app.models.enums import TaskStatus, VersionType
        from tests.g4 import builders as B
        v = B.version(db, 2029, VersionType.DRAFT, "2029 编制稿")
        db.add(BudgetTask(name="2029 年度预算编制", year=2029, deadline=date(2028, 12, 20), version_id=v.id,
                          status=TaskStatus.IN_PROGRESS))
        db.flush()
    elif name == "flow28":
        S.flow28(db)
    elif name == "flow29":
        from tests.acceptance.test_flow_29 import april, q1_history, setup_budget_and_rule
        setup_budget_and_rule(g3)
        g3.set_now("2027-04-20 10:00")
        g3.import_file(q1_history(g3), "kingdee-history")
        out = {"april": str(april(g3, 200000, "金蝶_2027-04_利润表.xlsx")),
               "corrected": str(april(g3, 208000, "金蝶_2027-04_利润表_更正.xlsx"))}
    elif name == "flow31":
        # 驱动假设「Amazon 美国站 11 月销量为 10 月的 1.6 倍」：v1 → v2（1.6 倍）已生效
        g3.set_now("2027-02-01 09:00")
        chen = g3.client("陈思雨")
        r = chen.post("/api/rules", json={"type": "driver_assumption",
                                          "name": "Amazon 美国站 11 月销量为 10 月的 1.6 倍",
                                          "content_text": "Amazon 美国站 11 月销量 = 10 月 × 1.4",
                                          "credibility": "unverified", "basis": "2025 年经验"})
        assert r.status_code == 200, r.text
        driver = r.json()
        g3.client("林晓").post(f"/api/approvals/rule-version/{driver['versions'][0]['id']}/approve")
        g3.set_now("2027-03-01 09:00")
        chen = g3.client("陈思雨")
        r = chen.post(f"/api/rules/{driver['id']}/versions",
                      json={"content_text": "Amazon 美国站 11 月销量 = 10 月 × 1.6", "credibility": "observed",
                            "basis": "2025–2026 金蝶"})
        g3.client("林晓").post(f"/api/approvals/rule-version/{r.json()['versions'][0]['id']}/approve")
        out = {"driver": driver["id"]}
    elif name == "flow32":
        _v, unit = S.flow32(db)
        out = {"unit": unit.id}
    elif name != "default":
        raise SystemExit(f"未知场景 {name}")
    clock.set_now(None)
    return out


def reset(name: str) -> dict:
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine

    from app.config import BACKEND_DIR
    from app.db import session_factory
    from scripts.seed import seed
    eng = create_engine(URL)
    with eng.begin() as conn:
        conn.exec_driver_sql("DROP SCHEMA IF EXISTS public CASCADE")
        conn.exec_driver_sql("CREATE SCHEMA public")
    eng.dispose()
    cfg = Config(str(BACKEND_DIR / "alembic.ini"))
    cfg.attributes["url"] = URL
    command.upgrade(cfg, "head")
    with session_factory(URL)() as db:
        seed(db)
        out = scenario(db, name)
        db.commit()
    return out


def add(name: str) -> dict:
    from app.db import session_factory
    with session_factory(URL)() as db:
        out = scenario(db, name)
        db.commit()
    return out


def serve() -> None:
    import uvicorn

    import app.main as m
    from app.core.clock import clock
    if os.environ.get("E3_STATIC_DIR"):
        m.STATIC_DIR = Path(os.environ["E3_STATIC_DIR"]).resolve()
    if os.environ.get("E3_NOW"):
        clock.set_now(_at(os.environ["E3_NOW"]))
    uvicorn.run(m.create_app(), host="127.0.0.1", port=int(os.environ.get("E3_PORT", "8013")), log_level="warning")


def xlsx(path: str) -> list:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        rows = [[{"v": (str(c.value) if c.value is not None else None), "f": c.number_format} for c in r]
                for r in ws.iter_rows()]
        out.append({"title": ws.title, "rows": rows})
    return out


def main(argv: list[str]) -> None:
    cmd = argv[0]
    if cmd == "serve":
        serve()
        return
    if cmd == "xlsx":
        print(json.dumps(xlsx(argv[1]), ensure_ascii=False))
        return
    fn = {"reset": reset, "add": add}[cmd]
    print(json.dumps(fn(argv[1] if len(argv) > 1 else "default"), ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:])
