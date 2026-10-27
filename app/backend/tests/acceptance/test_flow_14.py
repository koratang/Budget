"""FLOW-14 Excel 模板下载与上传（API 层验收）。期望值逐字取自 TDD验收契约.md FLOW-14：
王浩模板只含 Amazon-US-01 的维度组合与 7 个驱动列，已保存值（如 FLOW-11 的 SPU 10、单价 80.00）已预填；
蒋帆模板只含四项比例列，无收入、销量、利润列；上传后页面数据与数据库值均未变化，错误清单列出 3 条，各含行号、字段与原因；
修正后上传成功并刷新合计
"""
import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import PlatformUnit
from tests.acceptance.e2_common import Ctx

DRIVERS = ["SPU 数量", "单 SPU 平均销量", "平均单价（USD）", "折扣率（%）", "退款率（%）", "广告费率（%）", "佣金率（%）"]


def snapshot(db, store_id):
    db.expire_all()
    return sorted((u.id, u.spu_count, u.avg_sales, u.avg_price, u.discount_rate, u.revenue_cny)
                  for u in db.scalars(select(PlatformUnit).where(PlatformUnit.store_id == store_id)))


def download(client, url):
    r = client.get(f"{url}/template")
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def upload(client, url, wb):
    buf = io.BytesIO()
    wb.save(buf)
    return client.post(f"{url}/template-upload", files={"file": ("Amazon-US-01_模板.xlsx", buf.getvalue(),
                                                                  "application/vnd.ms-excel")})


def test_flow_14(seeded, client):
    db = seeded
    ctx = Ctx(db, client)
    task = ctx.publish()
    ctx.fill_flow11_drivers(task)
    unit_id = ctx.unit(task, "Amazon-US-01")
    url = f"/api/filling-units/{unit_id}"

    ctx.login("wanghao")
    wb = download(client, url)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == ["店铺", "站点", "品类", "新老品", "月份", *DRIVERS]
    assert {r[0] for r in rows[1:]} == {"Amazon-US-01"}
    hit = next(r for r in rows[1:] if (r[2], r[3], r[4]) == ("浴室柜", "新品", "2027-03"))
    assert hit[5] == "10" and hit[7] == "80.00"

    ctx.login("jiangfan")
    lwb = download(client, f"/api/filling-units/{ctx.unit(task, '物流比例（物流物控）')}")
    for name in lwb.sheetnames:
        header = [str(x) for x in next(lwb[name].iter_rows(values_only=True))]
        assert [x for x in header if "率" in x] == ["头程费率（%）", "尾程费率（%）", "仓储费率（%）", "关税率（%）"]
        assert not any(w in x for x in header for w in ("收入", "销量", "利润"))

    ctx.login("wanghao")
    store_id = next(u for u in client.get("/api/filling-units?mine=1").json()["items"] if u["id"] == unit_id)[
        "scope_object_id"]
    before, g0 = snapshot(db, store_id), ctx.grid(f"{url}/grid")
    ws["F5"], ws["I9"], ws["A12"] = "-1", "130", "Amazon-JP-01"
    ws["F3"], ws["G3"] = "7", "5"
    body = upload(client, url, wb).json()
    assert body["status"] == "failed" and body["written"] == 0
    assert sorted((e["row"], e["field"]) for e in body["errors"]) == [(5, "SPU 数量"), (9, "折扣率"), (12, "店铺")]
    assert all(e["message"] for e in body["errors"])
    assert snapshot(db, store_id) == before
    assert ctx.grid(f"{url}/grid")["rows"] == g0["rows"]

    ws["F5"], ws["I9"], ws["A12"] = "0", "0", "Amazon-US-01"
    assert upload(client, url, wb).json()["status"] == "imported"
    g1 = ctx.grid(f"{url}/grid")
    tot = lambda g: next(r for r in g["rows"] if r["key"] == "total|sales_qty")["values"]  # noqa: E731
    assert tot(g1) != tot(g0)
    assert snapshot(db, store_id) != before
