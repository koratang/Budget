"""RULE-01 金额精度（AC-F06-09、REQ-F06-01/02；DEC-005）。期望值逐字取自 TDD验收契约.md RULE-01 行：

明细显示 0.33、0.33、0.33，合计显示 1.00，差额 0.01 ≤ 0.01 × 3；1.005 显示 1.01（四舍五入，不是银行家舍入的 1.00）；
接口金额以字符串返回、数据库列为 NUMERIC，后端代码与依赖中金额路径无 float；伪造字段被忽略，派生值按服务端公式重算；
页面、接口、导出三处同值（页面层见 app/frontend/e2e/rule_01.spec.mjs）。
"""
from __future__ import annotations

import io
import json
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, text, update

from app.core.money import fmt_amount
from app.models import PlatformUnit
from tests.g2.helpers import category_id, grid, login, minimal_units, publish, pu_key, row, save, store, unit_by_name
from tests.g2.test_flow_11 import fill_flow11_drivers, preset_ratios

BACKEND = Path(__file__).resolve().parents[2]
CATS = ("框镜", "灯镜", "镜柜")


def test_rule_01_half_up_rounding():
    parts = [Decimal("0.334")] * 3
    assert [fmt_amount(p) for p in parts] == ["0.33", "0.33", "0.33"]
    total = sum(parts, Decimal(0))
    assert fmt_amount(total) == "1.00"
    diff = abs(Decimal(fmt_amount(total)) - sum(Decimal(fmt_amount(p)) for p in parts))
    assert diff == Decimal("0.01") and diff <= Decimal("0.01") * 3
    assert fmt_amount(Decimal("1.005")) == "1.01"
    assert round(1.005, 2) == 1.0  # 对照：float / 银行家舍入会得到 1.00


def _seed_precision_cells(db, unit_version_id):
    """前置：Amazon-US-01 2027-01 三个新品明细人民币收入 0.334；2027-02 浴室柜新品 1.005（未舍入值，直接入库）。"""
    s = store(db, "Amazon-US-01")
    base = (PlatformUnit.version_id == unit_version_id, PlatformUnit.store_id == s.id, PlatformUnit.is_new.is_(True))
    for cat in CATS:
        db.execute(update(PlatformUnit).where(*base, PlatformUnit.month == date(2027, 1, 1),
                                              PlatformUnit.category_id == category_id(db, cat))
                   .values(revenue_cny=Decimal("0.334")))
    db.execute(update(PlatformUnit).where(*base, PlatformUnit.month == date(2027, 2, 1),
                                          PlatformUnit.category_id == category_id(db, "浴室柜"))
               .values(revenue_cny=Decimal("1.005")))


def _export_values(client, unit_id) -> dict:
    r = client.get("/api/export", params={"view": "budget-platform", "filters": json.dumps({"unit_id": unit_id})})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    jan, feb = header.index("1 月"), header.index("2 月")
    out = {}
    for rr in rows[1:]:
        if rr[2] == "收入（人民币）":
            key = (rr[0], rr[1])
            out[key] = tuple(None if v is None or v == "—" else format(Decimal(str(v)).quantize(Decimal("0.01")), "f")
                             for v in (rr[jan], rr[feb]))
    return out


def test_rule_01_api_and_export_same_values(seeded, client):
    db = seeded
    task_id = publish(client, db, minimal_units(db)).json()["task"]["id"]
    unit = unit_by_name(db, task_id, "Amazon-US-01")
    _seed_precision_cells(db, unit.version_id)
    login(client, db, "wanghao")
    g = grid(client, unit.id)
    details = [row(g, pu_key(db, "Amazon-US-01", c, True, "revenue_cny"))["values"]["2027-01"] for c in CATS]
    total = row(g, "total|revenue_cny")["values"]["2027-01"]
    bath = row(g, pu_key(db, "Amazon-US-01", "浴室柜", True, "revenue_cny"))["values"]["2027-02"]
    assert details == ["0.33", "0.33", "0.33"] and total == "1.00" and bath == "1.01"
    assert all(isinstance(v, str) for r in g["rows"] for v in r["values"].values() if v is not None)
    exp = _export_values(client, unit.id)
    assert [exp[(c, "新品")][0] for c in CATS] == ["0.33", "0.33", "0.33"]
    assert exp[("店铺合计", "—")][0] == "1.00"
    assert exp[("浴室柜", "新品")][1] == "1.01"


def test_rule_01_numeric_columns_strings_and_forged_fields(seeded, client):
    db = seeded
    # 数据库：整个 schema 没有浮点列；金额列为 NUMERIC
    floats = db.execute(text("SELECT table_name, column_name FROM information_schema.columns WHERE table_schema = "
                             "'public' AND data_type IN ('real', 'double precision')")).all()
    assert floats == []
    money_cols = db.execute(text(
        "SELECT table_name || '.' || column_name, data_type FROM information_schema.columns WHERE table_schema = "
        "'public' AND (column_name ~ '(revenue|amount|profit|price|cost|rate|sales_local)')")).all()
    assert money_cols and all(t == "numeric" for _c, t in money_cols), [c for c, t in money_cols if t != "numeric"]
    assert type(db.scalar(text("SELECT 1.005::numeric"))) is Decimal  # 驱动层返回 Decimal

    task_id = publish(client, db, minimal_units(db)).json()["task"]["id"]
    unit_id = unit_by_name(db, task_id, "Amazon-US-01").id
    preset_ratios(client, db, task_id)
    g = fill_flow11_drivers(client, db, unit_id)
    for r in g["rows"]:
        assert all(v is None or isinstance(v, str) for v in r["values"].values())
    assert all(isinstance(v, str) for v in g["kpis"].values())
    key = lambda item: pu_key(db, "Amazon-US-01", "浴室柜", True, item)  # noqa: E731
    r = save(client, unit_id, [(key("revenue_cny"), "2027-03", "999999999"),
                               (key("gross_profit"), "2027-03", "888888888"), (key("spu_count"), "2027-03", "20")], g)
    assert r.status_code == 200, r.text
    assert {i["row_key"] for i in r.json()["ignored"]} == {key("revenue_cny"), key("gross_profit")}
    u = db.scalar(select(PlatformUnit).where(PlatformUnit.store_id == store(db, "Amazon-US-01").id,
                                             PlatformUnit.category_id == category_id(db, "浴室柜"),
                                             PlatformUnit.is_new.is_(True), PlatformUnit.month == date(2027, 3, 1)))
    db.refresh(u)
    assert fmt_amount(u.revenue_cny) == "1,230,336.00"  # 20 × 120 × 80 × 0.9 × 7.12
    assert fmt_amount(u.gross_profit) == "196,853.76"
    # 输入边界：即使客户端以 JSON 数字提交，也先按十进制文本（15 位有效数字）转 Decimal，库内无二进制误差
    g2 = r.json()["grid"]
    base = row(g2, key("avg_price"))["values"]["2027-03"]
    r = client.patch(f"/api/filling-units/{unit_id}/cells",
                     json=[{"row_key": key("avg_price"), "field": "2027-03", "value": 80.1, "base": base}])
    assert r.status_code == 200 and r.json()["errors"] == [], r.text
    db.refresh(u)
    assert u.avg_price == Decimal("80.1") and str(u.avg_price) in ("80.1", "80.10", "80.1000")
    assert fmt_amount(u.revenue_cny) == format(Decimal(20 * 120) * Decimal("80.1") * Decimal("0.9") * Decimal("7.12"),
                                               ",.2f")


# 允许出现 float 的位置：Excel 读取边界（xlsx 数值单元格以双精度存储，读取后立即按十进制文本转 Decimal）与非金额参数
FLOAT_BOUNDARY = {
    "app/services/importer.py": 2,      # cell_text / cell_decimal：Decimal(repr(v))
    "app/services/masterdata.py": 1,    # SKU 表格：整数日期序号 / 编码文本
    "app/services/filling.py": 2,       # parse_cell / 模板上传：format(raw, ".15g") 转十进制文本后走 Decimal 解析
    "app/dingtalk/gateway.py": 1,       # HTTP 超时秒数（非金额）
    "app/core/money.py": 4,             # 解析函数拒绝 float 的说明与判断
}


def test_rule_01_no_float_in_backend_money_path():
    found = {}
    for p in sorted((BACKEND / "app").rglob("*.py")):
        rel = str(p.relative_to(BACKEND))
        src = p.read_text(encoding="utf-8")
        n = len(re.findall(r"\bfloat\b", src))
        if n:
            found[rel] = n
        assert not re.search(r"\bfloat\s*\(", src), rel                  # 无 float() 转换
        assert not re.search(r"(Float|Double|REAL)\b\s*[(,)]", src), rel  # 无浮点列类型
    assert found == FLOAT_BOUNDARY
    # 边界处立即转 Decimal：Excel 浮点按 repr 十进制文本转换，不参与运算
    imp = (BACKEND / "app/services/importer.py").read_text(encoding="utf-8")
    assert "return Decimal(repr(v))" in imp
    # 依赖：SQLAlchemy Numeric 默认 asdecimal；requirements 无数值计算库（numpy / pandas）进入金额路径
    req = (BACKEND / "requirements.txt").read_text(encoding="utf-8").lower()
    assert "numpy" not in req and "pandas" not in req
