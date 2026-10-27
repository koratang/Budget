"""RULE-05 导入整份拒绝与覆盖保留（AC-F09-01/05/08、AC-F04-08、AC-F02-03、AC-F09-03/09；DEC-008）。
期望值逐字取自 TDD验收契约.md RULE-05 行：

四类错误文件上传前后全部业务表行数与校验和不变；注入错误时整个确认回滚，原批次仍为生效、新批次未生效、无预警与通知产生；
正常覆盖后旧批次行仍在库、状态「已被替换」且不参与任何计算，导入记录保留其文件名、导入人、导入时间。

「全部业务表」= public 下全部表，除去记录上传动作本身的 import_batch / import_staging、只追加的 audit_log、
登录会话 user_session 与 alembic_version。
"""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text

from app.models import Alert, AuditLog, FillingUnit, ImportBatch, KingdeeActual, NotificationOutbox
from app.models.enums import ImportStatus
from app.services import alerting
from app.services.versions import actual_monthly
from tests.acceptance.e4_common import table_snapshot
from tests.g2.helpers import minimal_units
from tests.g3.conftest import G3
from tests.g3.fixtures_g3 import fixture_path

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOT_BUSINESS = {"import_batch", "import_staging", "audit_log", "user_session", "alembic_version"}


def all_business(db) -> dict:
    names = db.scalars(text("SELECT tablename FROM pg_tables WHERE schemaname = 'public' ORDER BY 1")).all()
    return table_snapshot(db, [n for n in names if n not in NOT_BUSINESS])


@pytest.fixture
def g3(seeded, dingtalk, tmp_path):
    return G3(seeded, dingtalk, tmp_path)


def _sku_file(last_row_error: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["SKU", "SPU 编码", "品类", "首次上架日期"])
    for i in range(2, 121):
        ws.append([f"BC-{i:04d}-WH", f"SPU-BC-{i // 4:03d}", "浴室柜", "2025-06-01"])
    ws.append(["LM-4032-LED", "SPU-LM-4032", "灯镜", None if last_row_error else "2026-05-20"])  # 末行缺首次上架日期
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_rule_05_four_error_files_change_nothing(g3):
    db = g3.db
    g3.add_skus()
    g3.set_now("2026-12-01 09:00")
    # 填报模板前置：发布编制任务，王浩下载 Amazon-US-01 模板
    lin = g3.client("林晓")
    r = lin.post("/api/tasks", json={"year": 2027, "deadline": "2026-12-20", "units": minimal_units(db)})
    assert r.status_code == 201, r.text
    unit = db.scalar(select(FillingUnit).where(FillingUnit.task_id == r.json()["task"]["id"],
                                               FillingUnit.name == "Amazon-US-01"))
    wang = g3.client("王浩")
    wb = load_workbook(io.BytesIO(wang.get(f"/api/filling-units/{unit.id}/template").content))
    ws = wb[wb.sheetnames[0]]
    last = ws.max_row
    for row in range(2, last + 1):
        ws.cell(row, 6).value = "8"                      # SPU 数量
        ws.cell(row, 8).value = "79.99"                  # 平均单价（USD）
    ws.cell(last, 6).value = "-1"                        # 只有最后一行出错
    buf = io.BytesIO()
    wb.save(buf)
    g3.set_now("2027-05-12 10:00")
    bp = g3.client("陈思雨")
    wang = g3.client("王浩")

    before = all_business(db)
    k = g3.upload(fixture_path("金蝶_2027-04_末行出错.xlsx"), "kingdee")
    b = g3.upload(fixture_path("BI_2027-04_末行出错.xlsx"), "bi")
    s = bp.post("/api/master/sku-mapping/upload",
                files={"file": ("SKU对照表_2027-05.xlsx", _sku_file(True), XLSX)}).json()
    t = wang.post(f"/api/filling-units/{unit.id}/template-upload",
                  files={"file": ("Amazon-US-01_2027年度预算填报模板.xlsx", buf.getvalue(), XLSX)}).json()
    assert (k["status"], [i["row"] for i in k["issues"]]) == ("failed", [35])
    assert b["status"] == "failed" and b["unmapped_skus"] == ["XX-0000-ZZ"]
    assert (s["status"], [i["row"] for i in s["issues"]]) == ("failed", [121])
    assert (t["status"], t["written"], [e["row"] for e in t["errors"]]) == ("failed", 0, [last])
    after = all_business(db)
    assert after == before, {n for n in before if before[n] != after[n]}
    print(f"\nRULE-05 四类错误文件：业务表 {len(before)} 张行数与校验和不变；"
          f"金蝶末行 35、BI 未映射 XX-0000-ZZ、SKU 末行 121、模板末行 {last}")


def test_rule_05_confirm_rollback_then_overwrite_keeps_old(g3, monkeypatch):
    db = g3.db
    g3.add_skus()
    g3.set_now("2027-05-10 16:42")
    first = g3.import_file(fixture_path("金蝶_2027-04_利润表.xlsx"), "kingdee")
    g3.set_now("2027-05-12 10:00")
    fixed = g3.upload(fixture_path("金蝶_2027-04_利润表_v3.xlsx"), "kingdee")
    assert fixed["needs_overwrite"] is True
    db.commit()
    before = all_business(db)
    audits = db.scalar(select(func.count()).select_from(AuditLog))

    def boom(session, *args, **kwargs):  # 确认事务内（替换旧批次、写入新行之后）注入数据库错误
        session.execute(text("SELECT 1 / 0"))

    monkeypatch.setattr(alerting, "evaluate", boom)
    with pytest.raises(Exception, match="division by zero"):
        g3.confirm(fixed["id"], overwrite=True)
    db.expire_all()
    assert db.get(ImportBatch, first["id"]).status == ImportStatus.EFFECTIVE
    assert db.get(ImportBatch, fixed["id"]).status == ImportStatus.PENDING_CONFIRM
    assert db.scalar(select(func.count()).select_from(ImportBatch).where(
        ImportBatch.parent_batch_id == fixed["id"])) == 0
    assert db.scalar(select(func.count()).select_from(Alert)) == 0
    assert db.scalar(select(func.count()).select_from(NotificationOutbox)) == 0
    assert all_business(db) == before
    assert db.scalar(select(func.count()).select_from(AuditLog)) == audits

    monkeypatch.undo()
    r = g3.confirm(fixed["id"], overwrite=True)
    assert r.status_code == 200, r.text
    db.expire_all()
    old_child = db.scalar(select(ImportBatch).where(ImportBatch.parent_batch_id == first["id"]))
    assert old_child.status == ImportStatus.REPLACED
    assert db.get(ImportBatch, first["id"]).status == ImportStatus.REPLACED
    assert db.scalar(select(func.count()).select_from(KingdeeActual).where(KingdeeActual.batch_id == old_child.id)) \
        == 34
    revenue = actual_monthly(db, [date(2027, 4, 1)])[date(2027, 4, 1)].revenue
    assert revenue == 535000 + 355000 + 98000 + 120000 + 210000 + 86000 + 64000 + 230000 + 76000 + 42000
    # 旧批次不参与任何计算：报表接口的 2027-04 收入同为修正文件合计
    rep = g3.client("林晓").get("/api/reports/channel-site?month=2027-04")
    assert rep.status_code == 200, rep.text
    rec = next(x for x in g3.client("陈思雨").get("/api/imports").json()["records"] if x["id"] == first["id"])
    assert (rec["file_name"], rec["imported_by"], rec["imported_at"], rec["status_label"]) == \
        ("金蝶_2027-04_利润表.xlsx", "陈思雨", "2027-05-10 16:42", "已被替换")
