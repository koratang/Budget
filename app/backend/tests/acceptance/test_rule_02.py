"""RULE-02 越权（NFR-003、REQ-F01-03、AC-F01-11、AC-F14-01；DEC-009）。期望值逐字取自 TDD验收契约.md RULE-02 行：

除登录与健康检查外每个路由都声明认证与权限依赖，新增未声明路由使测试失败；全部越权请求 100% 返回 403，
响应体仅为 {"code":"forbidden","message":"无权限"}，不含任何记录字段；列表与导出在服务端过滤，前端不收到范围外行；
页面直链显示「无权限」（浏览器层见 app/frontend/e2e/rule_02.spec.mjs）。

四类受限角色：王浩（运营，Amazon-US-01）、孙婷（B端，Ferguson / Menards）、吴静（部门负责人，人力资源部）、
蒋帆（物流物控）。
"""
from __future__ import annotations

import io
import json
import re
from datetime import date

import pytest
from fastapi import APIRouter, Depends
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.clock import clock
from app.core.scope import get_scope
from app.main import create_app
from app.models import Alert, BCustomer, Channel, Department, FillingUnit, Rule, RuleVersion, Store
from app.models.enums import VersionType
from tests.acceptance.e4_common import FORBIDDEN_BODY, client_as, iter_api_routes, route_guard
from tests.conftest import make_settings
from tests.g2.helpers import add_user, login
from tests.g3.conftest import G3, at
from tests.g4 import builders as B

ROLES = {"wanghao": "王浩", "sunting": "孙婷", "wujing": "吴静", "jiangfan": "蒋帆"}
ROLE_CODES = {"wanghao": {"operations"}, "sunting": {"b2b"}, "wujing": {"dept_head"}, "jiangfan": {"logistics"}}

# 「登录与健康检查」：无需登录即可访问的路由（含开发身份登录、未知 /api 路径的 404 兜底与 SPA 静态托管）
PUBLIC_ROUTES = {
    ("GET", "/api/health"), ("GET", "/api/auth/config"), ("GET", "/api/auth/dingtalk/qr-url"),
    ("GET", "/api/auth/dingtalk/callback"), ("POST", "/api/auth/dingtalk/h5"), ("POST", "/api/auth/logout"),
    ("GET", "/api/auth/session"), ("GET", "/api/auth/dev-users"), ("POST", "/api/auth/dev-login"),
    ("POST", "/api/dev/login-as"),
}
FALLBACK_ROUTES = {"/api/{rest:path}", "/{full_path:path}"}

# 只返回「本人」数据的数据范围路由：以服务端过滤断言（见 test_lists_and_exports_filtered_server_side）
SELF_SCOPED = {("GET", "/api/me"), ("GET", "/api/workbench"), ("GET", "/api/filling-units"),
               ("GET", "/api/dept-expense/my-unit"), ("GET", "/api/versions"), ("GET", "/api/versions/compare"),
               ("GET", "/api/alerts"), ("GET", "/api/versions/rolling/{version_id}")}


def _routes(app):
    out = []
    for r in iter_api_routes(app):
        for m in sorted(r.methods - {"HEAD"}):
            out.append((m, r.path, r))
    return out


def undeclared_routes(app) -> list[str]:
    """路由权限声明检查：除登录与健康检查外，每个路由都须有认证依赖，且有 require_roles 或 get_scope。"""
    bad = []
    for method, path, r in _routes(app):
        if path in FALLBACK_ROUTES or (method, path) in PUBLIC_ROUTES:
            continue
        auth, roles, scoped = route_guard(r)
        if not auth or (roles is None and not scoped):
            bad.append(f"{method} {path}")
    return bad


def test_rule_02_every_route_declares_auth_and_permission():
    app = create_app(make_settings())
    routes = _routes(app)
    assert len(routes) > 90  # 遍历的是完整注册表（含各功能组 include_router）
    assert undeclared_routes(app) == []


def test_rule_02_new_undeclared_route_fails_check():
    app = create_app(make_settings())
    probe = APIRouter()

    @probe.get("/probe/no-guard")
    def no_guard():
        return {"rows": []}

    @probe.get("/probe/with-scope")
    def with_scope(ctx=Depends(get_scope)):
        return {}

    app.include_router(probe, prefix="/api")
    assert undeclared_routes(app) == ["GET /api/probe/no-guard"]


# ---------------------------------------------------------------------------
# 场景数据
# ---------------------------------------------------------------------------
def _alert(db, object_type: str, obj, name: str) -> Alert:
    rv = db.scalar(select(RuleVersion).join(Rule, Rule.id == RuleVersion.rule_id).order_by(RuleVersion.id))
    a = Alert(rule_version_id=rv.id, object_type=object_type, object_id=str(obj.id), object_name=name,
              period=date(2027, 4, 1))
    db.add(a)
    db.flush()
    return a


@pytest.fixture
def world(seeded, dingtalk, tmp_path):
    db = seeded
    # 市场部补齐填报人与审核人，使「市场部费用」成为一个真实填报单位
    add_user(db, "liumin", "刘敏", "市场部", ["dept_budget"], {"department": [B.dept(db, "市场部").id]})
    add_user(db, "chenhao", "陈浩", "市场部", ["dept_head"], {"department": [B.dept(db, "市场部").id]}, head=True)
    clock.set_now(at("2026-12-01 09:00"))
    fin = client_as(db, dingtalk, "linxiao")
    units = [{"unit_type": "platform", "scope_object_id": B.store(db, s).id} for s in ("Amazon-US-01", "Amazon-DE-01")]
    units += [{"unit_type": "b2b", "scope_object_id": B.dept(db, d).id} for d in ("美国业务部", "外贸部")]
    units += [{"unit_type": "cost_ratio", "scope_object_id": None},
              {"unit_type": "logistics_ratio", "scope_object_id": None}]
    units += [{"unit_type": "dept_expense", "scope_object_id": B.dept(db, d).id} for d in ("人力资源部", "市场部")]
    r = fin.post("/api/tasks", json={"year": 2027, "deadline": "2026-12-20", "units": units})
    assert r.status_code == 201, r.text
    task_id = r.json()["task"]["id"]

    def unit(name):
        return db.scalar(select(FillingUnit).where(FillingUnit.task_id == task_id, FillingUnit.name.like(f"%{name}%")))

    w = {"task": task_id, "us": unit("Amazon-US-01"), "de": unit("Amazon-DE-01"), "b2b_us": unit("美国业务部"),
         "b2b_eu": unit("外贸部"), "logi": unit("物流比例"), "cost": unit("成本比例"), "hr": unit("人力资源部"),
         "mkt": unit("市场部")}
    assert all(w.values()), w
    w["alert_de"] = _alert(db, "store", B.store(db, "Amazon-DE-01"), "Amazon-DE-01")
    w["alert_us"] = _alert(db, "store", B.store(db, "Amazon-US-01"), "Amazon-US-01")
    w["alert_mkt"] = _alert(db, "department", B.dept(db, "市场部"), "市场部")
    w["clients"] = {u: client_as(db, dingtalk, u) for u in ROLES}
    w["db"] = db
    return w


def _ids(db):
    store = {s.name: s for s in db.scalars(select(Store))}
    cust = {c.name: c for c in db.scalars(select(BCustomer))}
    dept = {d.name: d for d in db.scalars(select(Department))}
    chan = {c.name: c for c in db.scalars(select(Channel))}
    return store, cust, dept, chan


def _pu(store, cat_id, item):
    return f"pu:{store.id}:{store.site_id}:{cat_id}:1|{item}"


def _check(failures, label, r):
    if r.status_code != 403 or r.json() != FORBIDDEN_BODY:
        failures.append(f"{label} → {r.status_code} {r.text[:200]}")


def _send(c, method, url, body=None, files=None):
    if files is not None:
        return c.request(method, url, files=files)
    if method in ("GET",):
        return c.get(url)
    return c.request(method, url, json=body if body is not None else {})


# ---------------------------------------------------------------------------
# 1. 功能级：遍历注册表中全部 require_roles 路由，四类角色凡不在声明角色内一律 403
# ---------------------------------------------------------------------------
def test_rule_02_role_gated_routes_forbidden(world):
    db = world["db"]
    params = {"unit_id": world["de"].id, "task_id": world["task"], "store_id": B.store(db, "Amazon-DE-01").id,
              "channel_id": 1, "customer_id": B.customer(db, "Bauhaus").id, "department_id": B.dept(db, "市场部").id,
              "account_id": 1, "user_id": 1, "batch_id": 1, "version_id": 1, "rule_id": 1, "reserve_id": 1,
              "adjustment_id": 1, "obj_type": "unit", "obj_id": world["mkt"].id, "alert_id": world["alert_de"].id,
              "unit_key": "platform:1", "key": "policy"}
    failures, checked = [], 0
    for method, path, r in _routes(world["clients"]["wanghao"].app):
        _auth, roles, _scoped = route_guard(r)
        if roles is None:
            continue
        url = re.sub(r"\{(\w+)(?::\w+)?\}", lambda m: str(params[m.group(1)]), path)
        for userid in ROLES:
            if ROLE_CODES[userid] & roles:
                continue
            c = world["clients"][userid]
            checked += 1
            _check(failures, f"{ROLES[userid]} {method} {url}", _send(c, method, url))
    print(f"\nRULE-02 功能级越权请求 {checked} 次")
    assert checked > 150
    assert failures == [], "\n".join(failures)


# ---------------------------------------------------------------------------
# 2. 数据范围：范围外的详情、列表筛选、导出、写入（含契约举例）
# ---------------------------------------------------------------------------
def scope_cases(world) -> list[tuple[str, str, str, str, dict | list | None]]:
    db = world["db"]
    store, cust, dept, chan = _ids(db)
    de, us, mkt, eu, b2b_us, logi = (world[k].id for k in ("de", "us", "mkt", "b2b_eu", "b2b_us", "logi"))
    cat = B.category(db, "浴室柜").id
    de_s, us_s = store["Amazon-DE-01"], store["Amazon-US-01"]
    task, a_de, a_mkt = world["task"], world["alert_de"].id, world["alert_mkt"].id
    ver = db.get(FillingUnit, us).version_id
    exp = "/api/export?view={}&filters={}".format
    cell = lambda key, m="2027-03", v="1": [{"row_key": key, "field": m, "value": v, "base": None}]  # noqa: E731
    return [
        # 王浩：请求 Amazon-DE-01 填报单位（详情 / 写入 / 导出）及其他范围外对象
        ("wanghao", "GET", f"/api/filling-units/{de}/grid", "detail", None),
        ("wanghao", "GET", f"/api/filling-units/{de}/template", "export", None),
        ("wanghao", "PATCH", f"/api/filling-units/{de}/cells", "write", cell(_pu(de_s, cat, "spu_count"))),
        ("wanghao", "PATCH", f"/api/filling-units/{us}/cells", "write", cell(_pu(de_s, cat, "spu_count"))),
        ("wanghao", "POST", f"/api/filling-units/{de}/fill-year", "write",
         {"row_key": _pu(de_s, cat, "avg_price"), "value": "80"}),
        ("wanghao", "POST", f"/api/filling-units/{de}/submit", "write", {}),
        ("wanghao", "PUT", f"/api/filling-units/{de}/over-limit-reasons", "write", {"reasons": {"2027-03": "测试"}}),
        ("wanghao", "POST", f"/api/filling-units/{de}/template-upload", "write", "FILE"),
        ("wanghao", "GET", exp("budget-platform", json.dumps({"unit_id": de})), "export", None),
        ("wanghao", "GET", f"/api/reports/channel-site?channel={chan['Wayfair'].id}", "list", None),
        ("wanghao", "GET", exp("report-channel", json.dumps({"channel": str(chan["Wayfair"].id)})), "export", None),
        ("wanghao", "GET", f"/api/reports/dept-expense?department_id={dept['市场部'].id}", "list", None),
        ("wanghao", "GET", "/api/reports/company", "list", None),
        ("wanghao", "GET", f"/api/dept-expense/{mkt}/grid", "detail", None),
        ("wanghao", "GET", f"/api/reserves?department_id={dept['市场部'].id}", "list", None),
        ("wanghao", "POST", "/api/reserves", "write", {"department_id": dept["市场部"].id, "name": "展会",
                                                      "amount": "10000", "plan_month": "2027-05"}),
        ("wanghao", "GET", "/api/tasks", "list", None),
        ("wanghao", "GET", f"/api/tasks/{task}/progress", "detail", None),
        ("wanghao", "GET", f"/api/alerts/{a_de}", "detail", None),
        ("wanghao", "POST", f"/api/alerts/{a_de}/resolve", "write", {"note": "已处理"}),
        ("wanghao", "GET", f"/api/versions/{ver}/summary?store_id={de_s.id}", "detail", None),
        ("wanghao", "GET", f"/api/versions/{ver}/summary", "detail", None),
        # 孙婷：请求 Bauhaus（外贸部）偏差与填报
        ("sunting", "GET", f"/api/filling-units/{eu}/grid", "detail", None),
        ("sunting", "PATCH", f"/api/filling-units/{b2b_us}/cells", "write",
         cell(f"b2b:{cust['Bauhaus'].id}|sales_local")),
        ("sunting", "GET", f"/api/reports/channel-site?channel={chan['Amazon'].id}", "list", None),
        ("sunting", "GET", exp("report-channel", json.dumps({"channel": str(chan["Amazon"].id)})), "export", None),
        ("sunting", "GET", exp("budget-platform", json.dumps({"unit_id": us})), "export", None),
        ("sunting", "GET", "/api/reports/category", "list", None),
        ("sunting", "GET", f"/api/alerts/{a_de}", "detail", None),
        # 吴静：请求市场部费用
        ("wujing", "GET", f"/api/reports/dept-expense?department_id={dept['市场部'].id}", "list", None),
        ("wujing", "GET", exp("report-dept", json.dumps({"department_id": dept["市场部"].id})), "export", None),
        ("wujing", "GET", f"/api/dept-expense/{mkt}/grid", "detail", None),
        ("wujing", "GET", f"/api/filling-units/{mkt}/template", "export", None),
        ("wujing", "PATCH", f"/api/filling-units/{mkt}/cells", "write",
         cell(f"de:{B.account(db, '差旅费').id}|amount")),
        ("wujing", "GET", f"/api/reserves?department_id={dept['市场部'].id}", "list", None),
        ("wujing", "GET", f"/api/approvals/unit/{mkt}", "detail", None),
        ("wujing", "POST", f"/api/approvals/unit/{mkt}/approve", "write", {}),
        ("wujing", "POST", f"/api/approvals/unit/{mkt}/reject", "write", {"reason": "退回"}),
        ("wujing", "GET", f"/api/alerts/{a_mkt}", "detail", None),
        ("wujing", "POST", f"/api/alerts/{a_mkt}/resolve", "write", {"note": "已处理"}),
        ("wujing", "GET", f"/api/filling-units/{de}/grid", "detail", None),
        # 蒋帆：请求平台填报收入列
        ("jiangfan", "GET", f"/api/filling-units/{us}/grid", "detail", None),
        ("jiangfan", "PATCH", f"/api/filling-units/{us}/cells", "write", cell(_pu(us_s, cat, "revenue_cny"))),
        ("jiangfan", "PATCH", f"/api/filling-units/{logi}/cells", "write",
         cell(f"ccr:{chan['Amazon'].id}:{store['Amazon-US-01'].site_id}:{cat}|purchase_cost_rate")),
        ("jiangfan", "GET", exp("budget-platform", json.dumps({"unit_id": us})), "export", None),
        ("jiangfan", "GET", f"/api/filling-units/{us}/template", "export", None),
        ("jiangfan", "GET", "/api/reports/channel-site", "list", None),
        ("jiangfan", "GET", exp("report-channel", "{}"), "export", None),
        ("jiangfan", "GET", f"/api/versions/{ver}/summary?store_id={us_s.id}", "detail", None),
    ]


def _xlsx(db):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["店铺", "站点", "品类", "新品/老品", "月份", "SPU 数"])
    wb.active.append(["Amazon-DE-01", "DE", "浴室柜", "新品", "2027-03", 5])
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": ("Amazon-DE-01_模板.xlsx", buf.getvalue(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_rule_02_scope_matrix_forbidden(world):
    db = world["db"]
    failures = []
    cases = scope_cases(world)
    for userid, method, url, _kind, body in cases:
        c = world["clients"][userid]
        r = _send(c, method, url, files=_xlsx(db)) if body == "FILE" else _send(c, method, url, body)
        _check(failures, f"{ROLES[userid]} {method} {url}", r)
    assert {k for *_x, k, _b in cases} == {"detail", "list", "export", "write"}
    assert failures == [], "\n".join(failures)


# ---------------------------------------------------------------------------
# 3. 滚动预测与目标调整中的范围外单位
# ---------------------------------------------------------------------------
@pytest.fixture
def rolling_world(seeded, dingtalk, tmp_path):
    g = G3(seeded, dingtalk, tmp_path)
    g.set_now("2027-05-12 10:00")
    stores = {s: {m: (600000, 150000) for m in range(1, 13)} for s in ("Amazon-US-01", "Amazon-DE-01")}
    g.make_version(stores=stores, depts={"人力资源部": 600000})
    fin = g.client("林晓")
    r = fin.post("/api/versions/rolling")
    assert r.status_code == 200, r.text
    rolling = r.json()
    r = fin.post("/api/versions/adjustments")
    assert r.status_code == 200, r.text
    adj = r.json()
    return g, rolling, adj


def test_rule_02_rolling_and_adjustment_scope(rolling_world, dingtalk):
    g, rolling, adj = rolling_world
    db = g.db
    de, us = B.store(db, "Amazon-DE-01"), B.store(db, "Amazon-US-01")
    cat = B.category(db, "浴室柜").id
    v, key_de = rolling["id"], f"platform:{de.id}"
    de_adj = db.scalar(select(FillingUnit).where(FillingUnit.version_id == adj["draft_version_id"],
                                                 FillingUnit.unit_type == "platform",
                                                 FillingUnit.scope_object_id == de.id))
    wang, jiang = client_as(db, dingtalk, "wanghao"), client_as(db, dingtalk, "jiangfan")
    wu = client_as(db, dingtalk, "wujing")
    cell = [{"row_key": f"pu:{de.id}:{de.site_id}:{cat}:0|avg_price", "field": "2027-06", "value": "1", "base": None}]
    cases = [
        (wang, "GET", f"/api/filling-units/rolling/{v}/{key_de}/grid", None),
        (wang, "PATCH", f"/api/filling-units/rolling/{v}/{key_de}/cells", cell),
        (wang, "POST", f"/api/filling-units/rolling/{v}/{key_de}/fill-year",
         {"row_key": cell[0]["row_key"], "value": "1"}),
        (wang, "POST", f"/api/filling-units/rolling/{v}/{key_de}/submit", {}),
        (wang, "POST", f"/api/versions/rolling/{v}/units/{key_de}/submit", {}),
        (wang, "GET", f"/api/filling-units/{de_adj.id}/grid", None),
        (wang, "POST", f"/api/versions/adjustments/units/{de_adj.id}/submit", {"adjust_reason": "调整"}),
        (wu, "GET", f"/api/versions/adjustments/{adj['id']}", None),
        (jiang, "GET", f"/api/filling-units/rolling/{v}/platform:{us.id}/grid", None),
    ]
    failures = []
    for c, method, url, body in cases:
        _check(failures, f"{method} {url}", _send(c, method, url, body))
    assert failures == [], "\n".join(failures)
    # 范围内的同类请求可用（证明 403 来自范围判断而非接口不可用）
    assert wang.get(f"/api/filling-units/rolling/{v}/platform:{us.id}/grid").status_code == 200
    card = wang.get(f"/api/versions/rolling/{v}").json()
    assert [u["name"] for u in card["units"]] == ["Amazon-US-01"]


# ---------------------------------------------------------------------------
# 4. 列表与导出在服务端过滤：前端收不到范围外行
# ---------------------------------------------------------------------------
def _xlsx_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    return "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v)


def test_rule_02_lists_and_exports_filtered_server_side(world):
    db = world["db"]
    c = world["clients"]
    # 实际与预算：Amazon-US-01 / Amazon-DE-01、Ferguson / Bauhaus、人力资源部 / 市场部
    orig = B.version(db, 2026, VersionType.ORIGINAL_TARGET, "V1 年初原始目标")
    for s in ("Amazon-US-01", "Amazon-DE-01"):
        B.platform_unit(db, orig, s, "2026-04", "浴室柜", True, spu=10, avg_sales="120", avg_price="80.00",
                        discount="0.10", ad_rate="0.12")
    for cu in ("Ferguson", "Bauhaus"):
        B.b2b_unit(db, orig, cu, "2026-04", sales_local="100000.00", refund_rate="0.03", purchase_cost_rate="0.40")
    for d in ("人力资源部", "市场部"):
        B.dept_expense(db, orig, d, "差旅费", "2026-04", "30000.00")
    kd = B.batch(db, "kingdee", "2026-04", "金蝶_2026-04_利润表.xlsx")
    for s, amt in (("Amazon-US-01", "554400.00"), ("Amazon-DE-01", "321000.00")):
        B.kingdee(db, kd, "2026-04", "平台销售收入", amt, store_name=s)
    for cu, amt in (("Ferguson", "700000.00"), ("Bauhaus", "456000.00")):
        B.kingdee(db, kd, "2026-04", "平台销售收入", amt, customer_name=cu)
    for d, amt in (("人力资源部", "28000.00"), ("市场部", "41000.00")):
        B.kingdee(db, kd, "2026-04", "差旅费", amt, dept_name=d)
    B.fx_actual(db, "2026-04", "USD", "7.2000")
    B.fx_actual(db, "2026-04", "EUR", "7.8000")

    # 王浩：填报单位列表、预警列表、渠道偏差与导出只含 Amazon-US-01
    units = c["wanghao"].get("/api/filling-units").json()["items"]
    assert {u["name"] for u in units} == {"Amazon-US-01"}
    alerts = c["wanghao"].get("/api/alerts").json()
    assert [a["object_name"] for a in alerts["items"]] == ["Amazon-US-01"]
    rep = c["wanghao"].get("/api/reports/channel-site?month=2026-04")
    assert rep.status_code == 200, rep.text
    body = rep.text
    assert "Amazon-US-01" in body and "Amazon-DE-01" not in body and "Bauhaus" not in body and "Ferguson" not in body
    ex = c["wanghao"].get("/api/export", params={"view": "report-channel", "filters": json.dumps({"month": "2026-04"})})
    assert ex.status_code == 200
    text = _xlsx_text(ex.content)
    assert "Amazon-US-01" in text and "Amazon-DE-01" not in text and "Bauhaus" not in text
    wayfair = db.scalar(select(Channel.id).where(Channel.name == "Wayfair"))
    r = c["wanghao"].get(f"/api/reports/channel-site?month=2026-04&channel={wayfair}")
    assert (r.status_code, r.json()) == (403, FORBIDDEN_BODY)
    r = c["wanghao"].get("/api/export", params={"view": "report-channel",
                                                 "filters": json.dumps({"month": "2026-04", "channel": str(wayfair)})})
    assert (r.status_code, r.json()) == (403, FORBIDDEN_BODY)

    # 孙婷：渠道偏差 B 端只含 Ferguson（Bauhaus 范围外）
    rep = c["sunting"].get("/api/reports/channel-site?month=2026-04")
    assert rep.status_code == 200, rep.text
    assert "Ferguson" in rep.text and "Bauhaus" not in rep.text and "Amazon-US-01" not in rep.text
    ex = c["sunting"].get("/api/export", params={"view": "report-channel", "filters": json.dumps({"month": "2026-04"})})
    text = _xlsx_text(ex.content)
    assert "Ferguson" in text and "Bauhaus" not in text

    # 吴静：部门费用报告与导出只含人力资源部
    rep = c["wujing"].get("/api/reports/dept-expense?year=2026")
    assert rep.status_code == 200, rep.text
    assert [d["name"] for d in rep.json()["departments"]] == ["人力资源部"]
    assert "市场部" not in rep.text
    ex = c["wujing"].get("/api/export", params={"view": "report-dept", "filters": json.dumps({"year": 2026})})
    text = _xlsx_text(ex.content)
    assert "人力资源部" in text and "市场部" not in text
    alerts = c["wujing"].get("/api/alerts").json()
    assert alerts["items"] == []

    # 蒋帆：物流比例单位的表格不下发收入 / 利润列
    units = c["jiangfan"].get("/api/filling-units").json()["items"]
    assert [u["name"] for u in units] == ["物流比例（物流物控）"]
    g = c["jiangfan"].get(f"/api/filling-units/{world['logi'].id}/grid").json()
    keys = json.dumps(g, ensure_ascii=False)
    for word in ("revenue", "gross_profit", "收入", "毛利", "purchase_cost_rate"):
        assert word not in keys, word

    # 工作台与 /me 只含本人范围
    me = c["wanghao"].get("/api/me").json()
    assert me["scopes"]["stores"] == [B.store(db, "Amazon-US-01").id] and me["all_company"] is False
    wb = c["wanghao"].get("/api/workbench")
    assert wb.status_code == 200 and "Amazon-DE-01" not in wb.text
    assert c["wanghao"].get("/api/dept-expense/my-unit").status_code == 404
    vs = c["wanghao"].get("/api/versions")
    assert vs.status_code == 200
    cmp_ = c["wanghao"].get("/api/versions/compare")
    assert cmp_.status_code in (200, 404, 409, 422) and "Amazon-DE-01" not in cmp_.text


def test_rule_02_every_scoped_route_is_exercised(world):
    """全覆盖：注册表中每个数据范围路由（get_scope 且无 require_roles 限定）都出现在越权矩阵或本人数据过滤断言中。"""
    exercised = set(SELF_SCOPED)
    for _u, method, url, _k, _b in scope_cases(world):
        exercised.add((method, url.split("?")[0]))
    rolling = {("GET", "/api/filling-units/rolling/{version_id}/{unit_key}/grid"),
               ("PATCH", "/api/filling-units/rolling/{version_id}/{unit_key}/cells"),
               ("POST", "/api/filling-units/rolling/{version_id}/{unit_key}/fill-year"),
               ("POST", "/api/filling-units/rolling/{version_id}/{unit_key}/submit"),
               ("POST", "/api/versions/rolling/{version_id}/units/{unit_key}/submit"),
               ("POST", "/api/versions/adjustments/units/{unit_id}/submit"),
               ("GET", "/api/versions/adjustments/{adjustment_id}")}
    missing = []
    for method, path, r in _routes(world["clients"]["wanghao"].app):
        _auth, roles, scoped = route_guard(r)
        if not scoped or roles is not None or (method, path) in SELF_SCOPED or (method, path) in rolling:
            continue
        pattern = re.compile("^" + re.sub(r"\{\w+\}", r"[^/]+", path) + "$")
        if not any(m == method and pattern.match(u) for m, u in exercised):
            missing.append(f"{method} {path}")
    assert missing == []
